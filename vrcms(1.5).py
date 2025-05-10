import sys
print(sys)
from PyQt5 import QtCore, QtGui, QtWidgets
from PyQt5.QtWidgets import QProgressDialog, QApplication
from PyQt5.QtCore import Qt, QThread, pyqtSignal
from PyQt5.QtCore import QTimer
from PyQt5.QtWidgets import QMainWindow
from PyQt5.QtWidgets import QFileDialog
from PyQt5.QtGui import QMovie
from PyQt5.QtGui import QImage
from PyQt5.QtWidgets import QApplication, QPushButton
from PyQt5.QtGui import QPixmap, QScreen
from PyQt5.QtWidgets import *
from ui_mainwindow import Ui_MainWindow
from playsound import playsound
from itertools import repeat
from docx import Document
from docx.shared import Inches
from docx2pdf import convert
from datetime import datetime, timedelta
from dateutil import parser 
from contextlib import redirect_stdout, redirect_stderr
import docx
import webbrowser
import subprocess
import threading
import pandas as pd 
import select
import cv2
import serial
import time
from fpdf import FPDF
import os
import glob
from datetime import datetime
from datetime import date
import serial
import serial.tools.list_ports
import socket
import csv
import subprocess
import io
global host,port
host = '192.168.12.110'
port =  65000
global x
global p
global code
global m
global d
global lop_count
global i
global row
global image_path
global image_path1
global screen_name
global date_folder
global filtered_files
global rec_alarm,rec_alarm1,rec_alarm2,rec_alarm3
rec_alarm=0
rec_alarm1=0
rec_alarm2=0
rec_alarm3=0
global it
global main
it=0
lop_count=0
x=1
global comport
comport=2
code=0
record=0
global image_name
global current_datetime
global alarm_file
global vec_cont
vec_cont=1
global obsrv_count
obsrv_count=1
global recorded_count
recorded_count=1
url1  = "rtsp://admin:admin123@192.168.12.64:554/avstream/channel=1/stream=0.sdp"
url2 = "rtsp://admin:admin123@192.168.12.65:554/avstream/channel=1/stream=0.sdp"
#####################################################################################################################################
class VideoCaptureThread(QThread):
    # Signal to send frames to the main thread
    frame_received = pyqtSignal(QPixmap)
    connection_lost = pyqtSignal()
    def __init__(self, rtsp_url):
        super().__init__()
        self.rtsp_url = rtsp_url
        self.capture = None
        self.running = False
        self.connected=False
    def run(self):
        self.running = True
        while self.running:
            # Try to open the capture if not already open
            if not self.connected or not self.capture or not self.capture.isOpened():
                if self.capture:
                    self.capture.release()  # Ensure release before reconnecting
                self.capture = cv2.VideoCapture(self.rtsp_url)
                if self.capture.isOpened():
                    self.connected = True
                else:
                    self.connected = False
                    self.connection_lost.emit()
                    self.msleep(1000)  # Wait before retrying
            if self.capture and self.capture.isOpened():
                ret, frame = self.capture.read()
                if ret:
                    self.connected = True
                    # Convert the frame to QImage and emit it
                    videoFrame = QtGui.QImage(frame, frame.shape[1], frame.shape[0],
                                                       QtGui.QImage.Format_BGR888)
                    resize_video = videoFrame.scaled(640, 360)
                    convertFrame = QtGui.QPixmap.fromImage(resize_video)
                    self.frame_received.emit(convertFrame)
                    #print("j")
                else:
                    # Signal connection lost if no frame
                    self.connected=False
                    self.connection_lost.emit()
                    self.msleep(1000)  # Retry every second

    def stop(self):
        # Stop the thread and release resources
        self.running = False
        self.wait()
        if self.capture:
            self.capture.release()

########################################################################################################################
########################################################################################################################       
class MainWindow(QtWidgets.QMainWindow, Ui_MainWindow):
    def __init__(self):
        super().__init__()
        self.setupUi(self)
        self.pic = 0
        self.stop_event = threading.Event()
        self.stop_thread_event = threading.Event()
        self.thread = None  
        self.s = None
        self.capture_thread1 = VideoCaptureThread(url1) 
        # Connect signals to slots
        self.capture_thread1.frame_received.connect(self.display_frame)
        self.capture_thread1.connection_lost.connect(self.show_warning)
        self.capture_thread1.start()
        
        self.capture_thread2 = VideoCaptureThread(url2)    
        # Connect signals to slots
        self.capture_thread2.frame_received.connect(self.display_frame2)
        self.capture_thread2.connection_lost.connect(self.show_warning2)
        self.capture_thread2.start()
        
        #print("executing")
        ###################################################################################################################
        self.pushButton.clicked.connect(self.login)
        self.pushButton_8.clicked.connect(self.sblgin_cancel)
        self.pushButton_4.clicked.connect(self.login_view)
        self.pushButton_19.clicked.connect(self.sublogin_view)
        self.pushButton_2.clicked.connect(self.cancel)
        self.pushButton_8.clicked.connect(self.sblogin_cancel)
        self.pushButton_5.clicked.connect(self.live_settings)
        self.pushButton_45.clicked.connect(self.Back_live)
        self.pushButton_9.clicked.connect(self.sub_login)
        self.pushButton_10.clicked.connect(self.live_report)
        self.pushButton_90.clicked.connect(self.Quit)
        self.pushButton_34.clicked.connect(self.configuration)
        self.pushButton_35.clicked.connect(self.Alarm_level)
        self.pushButton_17.clicked.connect(self.config_set)
        self.pushButton_18.clicked.connect(self.config_get)
        self.pushButton_16.clicked.connect(self.alarm_set)
        self.pushButton_15.clicked.connect(self.alarm_get)
        self.pushButton_40.clicked.connect(self.communication_port)
        self.pushButton_21.clicked.connect(self.port_refresh)      
        self.pushButton_47.clicked.connect(self.port_connect)
        self.pushButton_12.clicked.connect(self.Serial_connect)
        self.pushButton_13.clicked.connect(self.Ethernet_connect)
        self.pushButton_42.clicked.connect(self.Bit)
        self.pushButton_55.clicked.connect(self.Bit)
        self.pushButton_44.clicked.connect(self.Bit)
        self.pushButton_37.clicked.connect(self.PBit)
        self.pushButton_38.clicked.connect(self.IBit)
        self.pushButton_27.clicked.connect(self.Home_Back1)
        self.pushButton_33.clicked.connect(self.Home_Back)
        self.pushButton_46.clicked.connect(self.Home_Back)
        self.pushButton_48.clicked.connect(self.Home_Back)
        self.pushButton_49.clicked.connect(self.Testmode_Back)
        self.pushButton_54.clicked.connect(self.back_to_report)
        self.pushButton_54.clicked.connect(self.back_to_report)
        self.pushButton_56.clicked.connect(self.Back_to_home)
        self.pushButton_26.clicked.connect(self.close_Back)
        self.pushButton_7.clicked.connect(self.open_pdf)
        self.pushButton_36.clicked.connect(self.Recorded_Alarm)
        self.pushButton_32.clicked.connect(self.save_data)
        self.pushButton_30.clicked.connect(self.Recorded_read)
        self.pushButton_41.clicked.connect(self.Generate_Report)
        self.pushButton_11.clicked.connect(self.record_image)
        self.pushButton_31.clicked.connect(self.Recorded_clear)
        self.pushButton_3.clicked.connect(self.genreport_open2)
        self.comboBox_5.activated.connect(self.report1)
        self.pushButton_43.clicked.connect(self.Testmode1)
        self.pushButton_89.clicked.connect(self.Testmode_again)
        self.tableWidget_5.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.tableWidget_5.itemSelectionChanged.connect(self.select_changed)
        self.pushButton_14.clicked.connect(self.load_csv_file)
        self.pushButton_6.clicked.connect(self.horn)
        self.serial_port = serial.Serial(timeout=1)
        self.timer5 = QtCore.QTimer()
        self.timer5.timeout.connect(self.datime_view)
        self.timer5.start(1000)
    def sblgin_cancel(self):
        self.stackedWidget.setCurrentIndex(12)
    def report1(self):
        if self.comboBox_5.currentText()=="Extracted Alarm Data":
            self.tableWidget_5.clear()
            self.pushButton_11.setStyleSheet("\n"
            "color: rgb(176,176,176);\n"
            "background-color: rgb(211, 211, 211);\n"
            "\n"
            "border-top: 3px solid black;\n"
            "border-right: 3px solid white;\n"
            "border-left: 3px solid black;\n"
            "border-bottom: 3px solid white;")
            self.pushButton_14.setStyleSheet("\n"
            "color: rgb(176,176,176);\n"
            "background-color: rgb(211, 211, 211);\n"
            "\n"
            "border-top: 3px solid black;\n"
            "border-right: 3px solid white;\n"
            "border-left: 3px solid black;\n"
            "border-bottom: 3px solid white;")
            self.timeEdit_2.setEnabled(False)
            self.timeEdit.setEnabled(False)
            self.pushButton_11.setEnabled(False)
            self.pushButton_14.setEnabled(False)
            column=["  "," "," "," "," "," "," "," "]
            self.tableWidget_5.setHorizontalHeaderLabels(column)
        elif self.comboBox_5.currentText()=="Recorded Alarm Data":
            self.tableWidget_5.clear()
            self.timeEdit_2.setEnabled(True)
            self.timeEdit.setEnabled(True)
            self.pushButton_11.setEnabled(False)
            self.pushButton_14.setEnabled(False)
            column=["  "," "," "," "," "," "," "," "]
            self.tableWidget_5.setHorizontalHeaderLabels(column)
        elif self.comboBox_5.currentText()=="Observed Data":
            self.pushButton_11.setStyleSheet("\n"
            "color: rgb(176,176,176);\n"
            "background-color: rgb(211, 211, 211);\n"
            "\n"
            "border-top: 3px solid black;\n"
            "border-right: 3px solid white;\n"
            "border-left: 3px solid black;\n"
            "border-bottom: 3px solid white;")
            self.pushButton_14.setStyleSheet("\n"
            "color: rgb(176,176,176);\n"
            "background-color: rgb(211, 211, 211);\n"
            "\n"
            "border-top: 3px solid black;\n"
            "border-right: 3px solid white;\n"
            "border-left: 3px solid black;\n"
            "border-bottom: 3px solid white;")
            self.tableWidget_5.clear()
            self.timeEdit_2.setEnabled(True)
            self.timeEdit.setEnabled(True)
            self.pushButton_11.setEnabled(False)
            self.pushButton_14.setEnabled(False)
            column=["  "," "," "," "," "," "," "," "]
            self.tableWidget_5.setHorizontalHeaderLabels(column)
        else:
            QMessageBox.critical(None, "ERROR", "Select Required Datatype")
    def genreport_open2(self):
        self.tableWidget_5.verticalScrollBar().setValue(0)
        st_filename = self.dateEdit.text()
        end_filename = self.dateEdit_2.text()
        start_time =st_filename+'/'+self.timeEdit.text()
        end_time = end_filename+'/'+self.timeEdit_2.text()
        if self.comboBox_5.currentText()=="Observed Data":#observed
            self.pushButton_11.setEnabled(False)
            self.pushButton_14.setEnabled(False)
            try:
                start_date_str = st_filename.split('.')[0]
                end_date_str = end_filename.split('.')[0]
                start_date = datetime.strptime(start_date_str, '%d-%m-%Y')
                end_date = datetime.strptime(end_date_str, '%d-%m-%Y')
            except ValueError as e:
                print(f"Error parsing dates: {e}")
                return
            path = os.getcwd()
            file_directory = os.path.join(path, 'Observed Data')
            if not os.path.exists(file_directory) or not os.listdir(file_directory):
                QtWidgets.QMessageBox.warning(self, "No Files", "The folder is empty or doesn't exist.")
                self.tableWidget_5.clear()
                column=[" "," "," "," "," "," "," "," "]
                self.tableWidget_5.setHorizontalHeaderLabels(column)
                return
            file_directory2 = os.listdir(file_directory)
            filtered_files = []
            for f in file_directory2:
                file_date_str = f.split('.')[0]
                try:
                    file_date = datetime.strptime(file_date_str, '%d-%m-%Y')
                    if start_date <= file_date <= end_date:
                        filtered_files.append(f)       
                except ValueError:
                    print(f"Skipping file with invalid date format: {f}")
            if not filtered_files:
                QtWidgets.QMessageBox.warning(self, "No Matching Files", "No files found within the specified date range.")
                return   
            self.tableWidget_5.clear()
            column=[" "," "," "," "," "," "," "," "]
            self.tableWidget_5.setHorizontalHeaderLabels(column)
            print("Filtered files within date range:", filtered_files)
            row_position = 0
            found = False
            for file in filtered_files:
                with open(os.path.join(file_directory, file), 'r') as csvfile:
                    reader = csv.reader(csvfile)
                    try:
                        header = next(reader)  
                    except StopIteration:
                        header = []
                    columns = ["Date","Time","Vehicle Count", "F1", "F2", "M1", "M2"]
                    self.tableWidget_5.setColumnCount(len(columns))
                    self.tableWidget_5.setHorizontalHeaderLabels(columns)
                    self.tableWidget_5.setStyleSheet("""
                QTableWidget {
                    background-color: lightgray;  /* Background color for the entire table */
                }
                QHeaderView::section {
                    background-color: darkgray;  /* Background color for the header */
                    font-weight: bold;           /* Make header text bold */
                }
                QTableWidget::item {
                    background-color: white;    /* Background color for table items */
                }
                QTableWidget::item:selected {
                    background-color: lightblue; /* Background color for selected items */
                }
            """)
                    header = self.tableWidget_5.horizontalHeader()
                    header.setDefaultSectionSize(280)  
                    header_font = QtGui.QFont()
                    header_font.setPointSize(17) 
                    header_font.setBold(True)     
                    header.setFont(header_font)
                    for row in reader:
                        if len(row) > 0:
                            current_time = row[1]
                            found = True
                            self.tableWidget_5.insertRow(row_position)
                            for col, value in enumerate(row):
                                item = QtWidgets.QTableWidgetItem(value)
                                data_font = item.font()
                                data_font.setPointSize(12)  
                                data_font.setBold(False)    
                                item.setFont(data_font)
                                item.setTextAlignment(QtCore.Qt.AlignCenter)
                                self.tableWidget_5.setItem(row_position, col, item)
                            row_position += 1
                                
            if not found:
                QtWidgets.QMessageBox.warning(self, "No Data", "No data found within the specified time range.")
                self.tableWidget_5.clear()
                column=[" "," "," "," "," "," "," "," "]
                self.tableWidget_5.setHorizontalHeaderLabels(column)
        elif self.comboBox_5.currentText()=="Recorded Alarm Data":#Recordeed
            try:
                self.pushButton_11.setEnabled(True)
                self.pushButton_14.setEnabled(True)
                start_date_str = st_filename.split('.')[0]
                end_date_str = end_filename.split('.')[0]
                start_date = datetime.strptime(start_date_str, '%d-%m-%Y')
                end_date = datetime.strptime(end_date_str, '%d-%m-%Y')
            except ValueError as e:
                print(f"Error parsing dates: {e}")
                return
            path = os.getcwd()
            file_directory = os.path.join(path, 'Recorded Alarm Data')
            if not os.path.exists(file_directory) or not os.listdir(file_directory):
                QtWidgets.QMessageBox.warning(self, "No Files", "The folder is empty or doesn't exist.")
                self.tableWidget_5.clear()
                column=[" "," "," "," "," "," "," "," "]
                self.tableWidget_5.setHorizontalHeaderLabels(column)
                return
            file_directory2 = os.listdir(file_directory)
            filtered_files = []
            for f in file_directory2:
                file_date_str = f.split('.')[0]
                try:
                    file_date = datetime.strptime(file_date_str, '%d-%m-%Y')
                    if start_date <= file_date <= end_date:
                        filtered_files.append(f)
                except ValueError:
                    print(f"Skipping file with invalid date format: {f}")
            if not filtered_files:
                QtWidgets.QMessageBox.warning(self, "No Matching Files", "No files found within the specified date range.")
                return
            self.tableWidget_5.clear()
            column=["  "," "," "," "," "," "," "," "]
            self.tableWidget_5.setHorizontalHeaderLabels(column)
            print("Filtered files within date range:", filtered_files)
            row_position = 0
            found = False
            for file in filtered_files:
                with open(os.path.join(file_directory, file), 'r') as csvfile:
                    reader = csv.reader(csvfile)
                    try:
                        header = next(reader)
                    except StopIteration:
                        header = []
                    self.pushButton_11.setStyleSheet("\n"
                    "color: rgb(0,0,0);\n"
                    "background-color: rgb(211, 211, 211);\n"
                    "border-top: 3px solid black;\n"
                    "border-right: 3px solid white;\n"
                    "border-left: 3px solid black;\n"
                    "border-bottom: 3px solid white;")
                    self.pushButton_14.setStyleSheet("\n"
                    "color: rgb(0,0,0);\n"
                    "background-color: rgb(211, 211, 211);\n"
                    "\n"
                    "border-top: 3px solid black;\n"
                    "border-right: 3px solid white;\n"
                    "border-left: 3px solid black;\n"
                    "border-bottom: 3px solid white;")
                    columns = ["Date","Time","Vehicle Count", "F1", "F2", "M1", "M2", "CAMERA-01", "CAMERA-02"]
                    self.tableWidget_5.setColumnCount(len(columns))
                    self.tableWidget_5.setHorizontalHeaderLabels(columns)
                    self.tableWidget_5.setStyleSheet("""
                    QTableWidget {
                        background-color: lightgray;  /* Background color for the entire table */
                    }
                    QHeaderView::section {
                        background-color: darkgray;  /* Background color for the header */
                        font-weight: bold;           /* Make header text bold */
                    }
                    QTableWidget::item {
                        background-color: white;    /* Background color for table items */
                    }
                    QTableWidget::item:selected {
                        background-color: lightblue; /* Background color for selected items */
                    }
                """)
                    header = self.tableWidget_5.horizontalHeader()
                    header.setDefaultSectionSize(280)  
                    header_font = QtGui.QFont()
                    header_font.setPointSize(17) 
                    header_font.setBold(True)     
                    header.setFont(header_font)
                    for row in reader:
                        if len(row) > 0:
                            current_time = row[0]
                            print(current_time)
                            found = True
                            self.tableWidget_5.insertRow(row_position)
                            for col, value in enumerate(row):
                                item = QtWidgets.QTableWidgetItem(value)
                                data_font = item.font()
                                data_font.setPointSize(12)  
                                data_font.setBold(False)    
                                item.setFont(data_font)
                                item.setTextAlignment(QtCore.Qt.AlignCenter)
                                self.tableWidget_5.setItem(row_position, col, item)
                            row_position += 1
            if not found:
                QtWidgets.QMessageBox.warning(self, "No Data", "No data found within the specified time range.")
                self.tableWidget_5.clear()
        elif self.comboBox_5.currentText()=="Extracted Alarm Data": #Extract 
            self.pushButton_11.setEnabled(False)
            self.pushButton_14.setEnabled(False)
            try:
                start_date_str = st_filename.split('.')[0]
                end_date_str = end_filename.split('.')[0]
                start_date = datetime.strptime(start_date_str, '%d-%m-%Y')
                end_date = datetime.strptime(end_date_str, '%d-%m-%Y')
                path = os.getcwd()
                file_directory = os.path.join(path, 'Extracted Alarm Data')
                file_directory2 = os.listdir(file_directory)
                filtered_files = []
                for f in file_directory2:
                    file_date_str = f.split('.')[0]
                    try:
                        file_date = datetime.strptime(file_date_str, '%d-%m-%Y')
                        if start_date <= file_date <= end_date:
                            filtered_files.append(f) 
                    except ValueError:
                        print(f"Skipping file with invalid date format: {f}")
                self.tableWidget_5.clear()
                column=[" "," "," "," "," "," "," "," "]
                self.tableWidget_5.setHorizontalHeaderLabels(column)
                print("Filtered files within date range:", filtered_files)
                row_position = 0
                found = False
                for file in filtered_files:
                    with open(os.path.join(file_directory, file), 'r') as csvfile:
                        reader = csv.reader(csvfile)
                        print(reader)
                        try:
                            print("s")
                            header = next(reader)  
                        except StopIteration:
                            print("s1")
                            header = []
                        columns = ["Date/Time","F1","F2","M1","M2"]
                        self.tableWidget_5.setColumnCount(len(columns))
                        self.tableWidget_5.setHorizontalHeaderLabels(columns)
                        self.tableWidget_5.setStyleSheet("""
                        QTableWidget {
                            background-color: lightgray;  /* Background color for the entire table */
                        }
                        QHeaderView::section {
                            background-color: darkgray;  /* Background color for the header */
                            font-weight: bold;           /* Make header text bold */
                        }
                        QTableWidget::item {
                            background-color: white;    /* Background color for table items */
                        }
                        QTableWidget::item:selected {
                            background-color: lightblue; /* Background color for selected items */
                        }
                    """)
                        header = self.tableWidget_5.horizontalHeader()
                        header.setDefaultSectionSize(280)  
                        header_font = QtGui.QFont()
                        header_font.setPointSize(17) 
                        header_font.setBold(True)     
                        header.setFont(header_font)
                        for row in reader:
                            if len(row)>0 and row[0] == start_time:  
                                found = True
                            self.tableWidget_5.insertRow(row_position)
                            for col, value in enumerate(row):
                                item = QtWidgets.QTableWidgetItem(value)
                                data_font = item.font()
                                data_font.setPointSize(12)  
                                data_font.setBold(False)    
                                item.setFont(data_font)
                                item.setTextAlignment(QtCore.Qt.AlignCenter)
                                self.tableWidget_5.setItem(row_position, col, item)
                            row_position += 1 
            except:
                column=[" "," "," "," "," "," "," "," "]
                self.tableWidget_5.setHorizontalHeaderLabels(column)
                self.stackedWidget.setCurrentIndex(11)
        else:
            QMessageBox.critical(None, "ERROR", "Select Required Datatype")
    def load_csv_file(self):
        st_filename = self.dateEdit.text()
        end_filename = self.dateEdit_2.text()
        start_time = f"{st_filename}/{self.timeEdit.text()}"
        end_time = f"{end_filename}/{self.timeEdit_2.text()}"
        start_date_str = st_filename.split('.')[0]
        end_date_str = end_filename.split('.')[0]
        try:
            start_date = datetime.strptime(start_date_str, '%d-%m-%Y')
            end_date = datetime.strptime(end_date_str, '%d-%m-%Y')
        except ValueError as e:
            print(f"Error parsing dates: {e}")
            return
        path = os.getcwd()
        file_directory = os.path.join(path, 'Recorded Alarm Data')
        filtered_files = self.get_filtered_files(file_directory, start_date, end_date)
        if not filtered_files:
            print("No files found in the specified date range.")
            return
        found = False
        for filename in filtered_files:
            file_path = os.path.join(file_directory, filename)
            found |= self.process_csv(file_path, start_time, end_time)
        if found:
            QMessageBox.information(None, 'Success', 'PDF documents created successfully!')
    def get_filtered_files(self, directory, start_date, end_date):
        filtered_files = []
        for filename in os.listdir(directory):
            file_date_str = filename.split('.')[0]
            try:
                file_date = datetime.strptime(file_date_str, '%d-%m-%Y')
                if start_date <= file_date <= end_date:
                    filtered_files.append(filename)
            except ValueError:
                print(f"Skipping file with invalid date format: {filename}")
        print("Filtered files within date range:", filtered_files)
        return filtered_files

    def process_csv(self, file_path, start_time, end_time):
        found = False
        try:
            with open(file_path, mode='r', newline='', encoding='utf-8') as file:
                reader = csv.reader(file)
                next(reader)
                for index, row in enumerate(reader):
                    #if row and start_time <= row[0] <= end_time:
                    found = True
                    self.create_pdf(row, index + 1, os.path.dirname(file_path))
        except Exception as e:
            print(f"Error processing CSV file: {e}")
        return found
    def create_pdf(self, row, index, output_dir):
        app = QApplication.instance()
        progress_dialog = QProgressDialog("Converting to PDF...", "Cancel", 0, 0)
        progress_dialog.setWindowTitle("Processing")
        progress_dialog.setWindowModality(Qt.WindowModal)
        progress_dialog.setValue(0)
        progress_dialog.resize(400, 100)
        progress_dialog.show()
        app.processEvents()
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=19)
        pdf.set_text_color(0, 102, 204)
        pdf.cell(200, 17, txt=f"Vehicle Radiation Contamination  Monitoring System", ln=True, align='C')
        pdf.set_font("Arial", size=12)
        pdf.set_text_color(0, 0, 0)
        headers = ['Date','Time','Vehicle_Count', 'F1', 'F2', 'M1', 'M2', 'Image_path1', 'Image_path2']
        for header, cell in zip(headers, row):
            pdf.cell(0, 10, txt=f'{header}: {cell}', ln=True)
            if header in ['Image_path1', 'Image_path2']:
                self.add_image(pdf, cell)
        current_date = datetime.now().strftime('%d-%m-%Y')
        output_directory = os.path.join(os.path.abspath(os.getcwd()), 'Pdf_File', current_date)
        os.makedirs(output_directory, exist_ok=True)
        date = row[0]
        vehicle_count = row[2]
        pdf_filename = os.path.join(output_directory, f"{date}_vehicle_{vehicle_count}.pdf")
        pdf.output(pdf_filename)
        print(f"Created PDF: {pdf_filename}")
        
    def add_image(self, pdf, image_path):
        try:
            if os.path.exists(image_path):
                pdf.image(image_path, w=200)
            else:
                pdf.cell(0, 10, txt=f'Image not found: {image_path}', ln=True, align='L')
        except Exception as e:
            raise Exception(f"Error adding image: {e}")

    def serial(self,comboBox_2):
        ports=serial.tools.list_ports.comports()
        for p in ports:
            self.comboBox_2.addItem(p.device)
    def Back_live(self):
        host = '192.168.12.110'
        port =  65000
        if comport==1:
             QMessageBox.information(None, "WARNING", f"Change to Ethernet communication") 
        elif comport==2:
             self.stackedWidget.setCurrentIndex(12)
             QtCore.QCoreApplication.processEvents()
             try:
                self.s.send("LD1!".encode())
                self.timer9 = QtCore.QTimer()
                self.timer9.timeout.connect(self.read_live_data)
                self.timer9.start(1000)
             except socket.error as e:
                self.label_224.setText("ETHERNET")
                self.label_224.setStyleSheet(
                    "background-color: rgb(218,0,0);\n"
                    "border-top: 6px solid black;\n"
                    "border-right: 6px solid white;\n"
                    "border-left: 6px solid black;\n"
                    "border-bottom: 6px solid white;"
                )
                self.timer9 = QtCore.QTimer()
                self.timer9.timeout.connect(self.read_live_data)
                self.timer9.start(1000)

    def port_connect2(self):
        device = self.comboBox_2.currentText()
        baudrate=115200
        host = '192.168.12.110'
        port =  65000
        def configure_serial_connection2():
            try:
                if hasattr(self, 's') and self.s:
                    self.s.close()
                if self.serial_port.isOpen():
                    self.serial_port.close()
                self.serial_port.port = device
                self.serial_port.baudrate = baudrate
                self.serial_port.timeout=1
                self.serial_port.open()
                print(f"Connected to {device}")
                self.label_224.setText('RS-485')
                self.label_224.setStyleSheet(
                    "background-color: rgb(62, 167, 46);\n"
                    "border-top: 6px solid black;\n"
                    "border-right: 6px solid white;\n"
                    "border-left: 6px solid black;\n"
                    "border-bottom: 6px solid white;"
                ) 
            except Exception as e:
                QMessageBox.information(None, "ERROR", f"Failed to connect to {device}: {e}")

        def configure_ethernet_connection2():
            try:
                if hasattr(self, 's') and self.s:
                    self.s.close()  
                if self.serial_port.isOpen():
                    self.serial_port.close()
                self.s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
                self.s.settimeout(1)
                self.s.connect((host, port))
                print("Ethernet connected2")
                self.label_224.setText("ETHERNET")
                self.label_224.setStyleSheet(
                    "background-color: rgb(62, 167, 46);\n"
                    "border-top: 6px solid black;\n"
                    "border-right: 6px solid white;\n"
                    "border-left: 6px solid black;\n"
                    "border-bottom: 6px solid white;"
                )
            except socket.error as e:
                self.label_224.setStyleSheet(
                    "background-color: rgb(218,0,0);\n"
                    "border-top: 6px solid black;\n"
                    "border-right: 6px solid white;\n"
                    "border-left: 6px solid black;\n"
                    "border-bottom: 6px solid white;"
                )
        if comport == 1:
            configure_serial_connection2()
        elif comport == 2:
            configure_ethernet_connection2()
        else:            
            self.label_224.setStyleSheet("background-color: rgb(218, 0, 0);")
    def port_connect3(self):
        host = '192.168.12.110'
        port =  65000
        global lop_count
        lop_count=0
        def configure_ethernet_connection():
            try:
                if hasattr(self, 's') and self.s:
                    self.s.close()  
                self.s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
                self.s.settimeout(5)
                self.s.connect((host, port))
                print("Ethernet connected3")
                self.label_224.setText("ETHERNET")
                self.label_224.setStyleSheet(
                    "background-color: rgb(62, 167, 46);\n"
                    "border-top: 6px solid black;\n"
                    "border-right: 6px solid white;\n"
                    "border-left: 6px solid black;\n"
                    "border-bottom: 6px solid white;"
                ) 
                self.s.send("LD1!".encode())
            except socket.error as e:
                pass
        if comport == 1:
            print("a")
        elif comport == 2:
            configure_ethernet_connection()
        else:
            self.label_224.setStyleSheet("background-color: rgb(218, 0, 0);")
    def port_connect(self):
        device = self.comboBox_2.currentText()
        baudrate=115200
        host = '192.168.12.110'
        port =  65000
        def configure_serial_connection():
            try:
                if hasattr(self, 's') and self.s:
                    self.s.close()
                if self.serial_port.isOpen():
                    self.serial_port.close()
                self.serial_port.port = device
                self.serial_port.baudrate = baudrate
                self.serial_port.timeout=1
                self.serial_port.open()
                print(f"Connected to {device}")
                self.stackedWidget.setCurrentIndex(1)
            except Exception as e:
                QMessageBox.information(None, "ERROR", f"Failed to connect to {device}: {e}")
                self.stackedWidget.setCurrentIndex(1) 
        def configure_ethernet_connection():
            try:
                if hasattr(self, 's') and self.s:
                    self.s.close()  
                self.s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
                self.s.settimeout(3)
                self.s.connect((host, port))
                self.label_224.setText("ETHERNET")
                self.label_224.setStyleSheet(
                    "background-color: rgb(62, 167, 46);\n"
                    "border-top: 6px solid black;\n"
                    "border-right: 6px solid white;\n"
                    "border-left: 6px solid black;\n"
                    "border-bottom: 6px solid white;"
                )
                self.stackedWidget.setCurrentIndex(1)
            except socket.error as e:
                QMessageBox.critical(self, "ERROR", f"Failed to connect to Ethernet: {e}")
                self.label_224.setStyleSheet(
                    "background-color: rgb(218,0,0);\n"
                    "border-top: 6px solid black;\n"
                    "border-right: 6px solid white;\n"
                    "border-left: 6px solid black;\n"
                    "border-bottom: 6px solid white;"
                )
                self.stackedWidget.setCurrentIndex(1)
        if comport == 1:
            configure_serial_connection()
        elif comport == 2:
            configure_ethernet_connection()
        else:
            self.stackedWidget.setCurrentIndex(1)
            self.label_224.setStyleSheet("background-color: rgb(218, 0, 0);")
    def Testmode1(self):
        try:
            if comport==1:
                device = self.comboBox_2.currentText()
                baudrate=115200
                host = '192.168.12.110'
                port =  65000
                if self.serial_port.isOpen():
                    self.serial_port.write(str("TM!").encode())
                    self.stackedWidget.setCurrentIndex(9)
                    QtCore.QCoreApplication.processEvents()
                    test_data=self.serial_port.readline()
                    print(test_data)
                    serialsize=len(test_data)
                    d1=test_data.decode('ascii')
                    if d1.startswith("TD1"):
                        pmtf1=d1[4]+d1[5]+d1[6]+d1[7]
                        self.label_147.setText(pmtf1)
                        pmt2f1=d1[9]+d1[10]+d1[11]+d1[12]
                        self.label_150.setText(pmt2f1)
                        pmttotal1=d1[14]+d1[15]+d1[16]+d1[17]
                        self.label_145.setText(pmttotal1)
                        pmtf2=d1[26]+d1[27]+d1[28]+d1[29]
                        self.label_143.setText(pmtf2)
                        pmt2f2=d1[31]+d1[32]+d1[33]+d1[34]
                        self.label_144.setText(pmt2f2)
                        pmttotal2=d1[36]+d1[37]+d1[38]+d1[39]
                        self.label_153.setText(pmttotal2)
                        m1=d1[48]+d1[49]+d1[50]+d1[51]
                        self.label_148.setText(m1)
                        pmt2m1=d1[53]+d1[54]+d1[55]+d1[56]
                        self.label_151.setText( pmt2m1)
                        pmttotal3=d1[58]+d1[59]+d1[60]+d1[61]
                        self.label_154.setText(pmttotal3)
                        pmtm2=d1[70]+d1[71]+d1[72]+d1[73]
                        self.label_149.setText(pmtm2)
                        pmt2m2=d1[75]+d1[76]+d1[77]+d1[78]
                        self.label_152.setText( pmt2m2)
                        pmttotal4=d1[80]+d1[81]+d1[82]+d1[83]
                        self.label_155.setText(pmttotal4)
                    elif d1.startswith("TD2"):
                        self.label_147.setText("")
                        self.label_150.setText("")
                        self.label_145.setText("")
                        pmtf2=d1[26]+d1[27]+d1[28]+d1[29]
                        self.label_143.setText(pmtf2)
                        pmt2f2=d1[31]+d1[32]+d1[33]+d1[34]
                        self.label_144.setText(pmt2f2)
                        pmttotal2=d1[36]+d1[37]+d1[38]+d1[39]
                        self.label_153.setText(pmttotal2)
                        m1=d1[48]+d1[49]+d1[50]+d1[51]
                        self.label_148.setText(m1)
                        pmt2m1=d1[53]+d1[54]+d1[55]+d1[56]
                        self.label_151.setText( pmt2m1)
                        pmttotal3=d1[58]+d1[59]+d1[60]+d1[61]
                        self.label_154.setText(pmttotal3)
                        pmtm2=d1[70]+d1[71]+d1[72]+d1[73]
                        self.label_149.setText(pmtm2)
                        pmt2m2=d1[75]+d1[76]+d1[77]+d1[78]
                        self.label_152.setText( pmt2m2)
                        pmttotal4=d1[80]+d1[81]+d1[82]+d1[83]
                        self.label_155.setText(pmttotal4)
                    elif d1.startswith("TD3"):
                        self.label_147.setText("")
                        self.label_150.setText("")
                        self.label_145.setText("")
                        self.label_143.serText("")
                        self.label_144.setText("")
                        self.label_153.setText("")
                        m1=d1[48]+d1[49]+d1[50]+d1[51]
                        self.label_148.setText(m1)
                        pmt2m1=d1[53]+d1[54]+d1[55]+d1[56]
                        self.label_151.setText( pmt2m1)
                        pmttotal3=d1[58]+d1[59]+d1[60]+d1[61]
                        self.label_154.setText(pmttotal3)
                        pmtm2=d1[70]+d1[71]+d1[72]+d1[73]
                        self.label_149.setText(pmtm2)
                        pmt2m2=d1[75]+d1[76]+d1[77]+d1[78]
                        self.label_152.setText(pmt2m2)
                        pmttotal4=d1[80]+d1[81]+d1[82]+d1[83]
                        self.label_155.setText(pmttotal4)
                    elif d1.startswith("TD4"):
                        self.label_147.setText("")
                        self.label_150.setText("")
                        self.label_145.setText("")
                        self.label_143.serText("")
                        self.label_144.setText("")
                        self.label_153.setText("")
                        self.label_148.setText("")
                        self.label_151.setText("")
                        self.label_154.setText("")
                        pmtm2=d1[70]+d1[71]+d1[72]+d1[73]
                        self.label_149.setText(pmtm2)
                        pmt2m2=d1[75]+d1[76]+d1[77]+d1[78]
                        self.label_152.setText( pmt2m2)
                        pmttotal4=d1[80]+d1[81]+d1[82]+d1[83]
                        self.label_155.setText(pmttotal4)
                    else:
                        print("Error")
                else:
                    self.stackedWidget.setCurrentIndex(9)
            elif comport==2:
                 QMessageBox.information(None, "WARNING", f"Change to RS485 communication")    
        except:
            self.stackedWidget.setCurrentIndex(9)

    def select_changed(self):
        global image_path
        selected_items = self.tableWidget_5.selectedItems()
        if selected_items:
            selected_row = selected_items[0].row()
            column_1_value = self.tableWidget_5.item(selected_row, 0).text() if self.tableWidget_5.item(selected_row, 0) else ""
            column_2_value = self.tableWidget_5.item(selected_row, 1).text() if self.tableWidget_5.item(selected_row, 1) else ""
            column_3_value = self.tableWidget_5.item(selected_row, 2).text() if self.tableWidget_5.item(selected_row, 2) else ""
            column_4_value = self.tableWidget_5.item(selected_row, 3).text() if self.tableWidget_5.item(selected_row, 3) else ""
            column_5_value = self.tableWidget_5.item(selected_row, 4).text() if self.tableWidget_5.item(selected_row, 4) else ""
            column_6_value = self.tableWidget_5.item(selected_row, 5).text() if self.tableWidget_5.item(selected_row, 5) else ""
            column_7_value = self.tableWidget_5.item(selected_row, 6).text() if self.tableWidget_5.item(selected_row, 6) else ""
            column_8_value = self.tableWidget_5.item(selected_row, 7).text() if self.tableWidget_5.item(selected_row, 7) else ""
            column_9_value = self.tableWidget_5.item(selected_row, 8).text() if self.tableWidget_5.item(selected_row, 8) else ""
            row_values = []
            for column in range(self.tableWidget_5.columnCount()):
                item = self.tableWidget_5.item(selected_row, column)
                if item:
                    row_values.append(item.text())
            print(f"Values in selected row: {row_values}")
            self.label_287.setText(f"{column_1_value}/{column_2_value}")
            self.label_288.setText(f"{column_3_value}")
            self.label_309.setText(f"{column_4_value}")
            self.label_292.setText(f"{column_5_value}")
            self.label_293.setText(f"{column_6_value}")
            self.label_294.setText(f"{column_7_value}")
            pixmap_1 = QPixmap(column_8_value)
            pixmap_2 = QPixmap(column_9_value)
            if not pixmap_1.isNull():
                pixmap_1 = pixmap_1.scaled(self.label_289.size(), aspectRatioMode=Qt.AspectRatioMode.KeepAspectRatio)
                self.label_289.setPixmap(pixmap_1)
            if not pixmap_2.isNull():
                pixmap_2 = pixmap_2.scaled(self.label_291.size(), aspectRatioMode=Qt.AspectRatioMode.KeepAspectRatio)
                self.label_291.setPixmap(pixmap_2)

    def Testmode_again(self):
       try:
            if comport==1:
                self.serial_port.close()
                device = self.comboBox_2.currentText()
                baudrate=115200
                host = '192.168.12.110'
                port =  65000
                self.port_connect2()
                if self.serial_port.isOpen():
                    self.serial_port.write(str("TM!").encode())
                    self.stackedWidget.setCurrentIndex(9)
                    QtCore.QCoreApplication.processEvents()
                    test_data=self.serial_port.readline()
                    print(test_data)
                    serialsize=len(test_data)
                    d1=test_data.decode('ascii')
                    if d1.startswith("TD1"):
                        pmtf1=d1[4]+d1[5]+d1[6]+d1[7]
                        self.label_147.setText(pmtf1)
                        pmt2f1=d1[9]+d1[10]+d1[11]+d1[12]
                        self.label_150.setText(pmt2f1)
                        pmttotal1=d1[14]+d1[15]+d1[16]+d1[17]
                        self.label_145.setText(pmttotal1)
                        pmtf2=d1[26]+d1[27]+d1[28]+d1[29]
                        self.label_143.setText(pmtf2)
                        pmt2f2=d1[31]+d1[32]+d1[33]+d1[34]
                        self.label_144.setText(pmt2f2)
                        pmttotal2=d1[36]+d1[37]+d1[38]+d1[39]
                        self.label_153.setText(pmttotal2)
                        m1=d1[48]+d1[49]+d1[50]+d1[51]
                        self.label_148.setText(m1)
                        pmt2m1=d1[53]+d1[54]+d1[55]+d1[56]
                        self.label_151.setText( pmt2m1)
                        pmttotal3=d1[58]+d1[59]+d1[60]+d1[61]
                        self.label_154.setText(pmttotal3)
                        pmtm2=d1[70]+d1[71]+d1[72]+d1[73]
                        self.label_149.setText(pmtm2)
                        pmt2m2=d1[75]+d1[76]+d1[77]+d1[78]
                        self.label_152.setText( pmt2m2)
                        pmttotal4=d1[80]+d1[81]+d1[82]+d1[83]
                        self.label_155.setText(pmttotal4)
                    elif d1.startswith("TD2"):
                        self.label_147.setText("")
                        self.label_150.setText("")
                        self.label_145.setText("")
                        pmtf2=d1[26]+d1[27]+d1[28]+d1[29]
                        self.label_143.setText(pmtf2)
                        pmt2f2=d1[31]+d1[32]+d1[33]+d1[34]
                        self.label_144.setText(pmt2f2)
                        pmttotal2=d1[36]+d1[37]+d1[38]+d1[39]
                        self.label_153.setText(pmttotal2)
                        m1=d1[48]+d1[49]+d1[50]+d1[51]
                        self.label_148.setText(m1)
                        pmt2m1=d1[53]+d1[54]+d1[55]+d1[56]
                        self.label_151.setText( pmt2m1)
                        pmttotal3=d1[58]+d1[59]+d1[60]+d1[61]
                        self.label_154.setText(pmttotal3)
                        pmtm2=d1[70]+d1[71]+d1[72]+d1[73]
                        self.label_149.setText(pmtm2)
                        pmt2m2=d1[75]+d1[76]+d1[77]+d1[78]
                        self.label_152.setText( pmt2m2)
                        pmttotal4=d1[80]+d1[81]+d1[82]+d1[83]
                        self.label_155.setText(pmttotal4)
                    elif d1.startswith("TD3"):
                        self.label_147.setText("")
                        self.label_150.setText("")
                        self.label_145.setText("")
                        self.label_143.serText("")
                        self.label_144.setText("")
                        self.label_153.setText("")
                        m1=d1[48]+d1[49]+d1[50]+d1[51]
                        self.label_148.setText(m1)
                        pmt2m1=d1[53]+d1[54]+d1[55]+d1[56]
                        self.label_151.setText( pmt2m1)
                        pmttotal3=d1[58]+d1[59]+d1[60]+d1[61]
                        self.label_154.setText(pmttotal3)
                        pmtm2=d1[70]+d1[71]+d1[72]+d1[73]
                        self.label_149.setText(pmtm2)
                        pmt2m2=d1[75]+d1[76]+d1[77]+d1[78]
                        self.label_152.setText(pmt2m2)
                        pmttotal4=d1[80]+d1[81]+d1[82]+d1[83]
                        self.label_155.setText(pmttotal4)
                    elif d1.startswith("TD4"):
                        self.label_147.setText("")
                        self.label_150.setText("")
                        self.label_145.setText("")
                        self.label_143.serText("")
                        self.label_144.setText("")
                        self.label_153.setText("")
                        self.label_148.setText("")
                        self.label_151.setText("")
                        self.label_154.setText("")
                        pmtm2=d1[70]+d1[71]+d1[72]+d1[73]
                        self.label_149.setText(pmtm2)
                        pmt2m2=d1[75]+d1[76]+d1[77]+d1[78]
                        self.label_152.setText( pmt2m2)
                        pmttotal4=d1[80]+d1[81]+d1[82]+d1[83]
                        self.label_155.setText(pmttotal4)
                    else:
                        print("Error") 
            elif comport==2:
                 QMessageBox.information(None, "WARNING", f"Change to RS485 communication")
            else:
                print("ERROR")
       except:
            self.stackedWidget.setCurrentIndex(9)

            
    def save_data(self):
        global it
        try:
            if comport==1:
                 QMessageBox.information(None, "WARNING", f"Change to Ethernet communication")  
            elif comport==2:
                try:
                    self.s.send("SD!".encode())
                    it = 0
                    global rec_alarm1
                    rec_alarm1=1
                    self.tableWidget_3.setEnabled(False)
                    self.gif_exec1()
                    threading.Thread(target=self.read_socket_data2).start()
                except Exception as e:
                    print("AAA")
                     
        except Exception as e:
            print(f"Error sending command to serial port: {e}")

            
    def gif_exec1(self):
        self.pushButton_46.setEnabled(False)
        self.pushButton_32.setEnabled(False)
        self.movie = QtGui.QMovie(":/newPrefix/newPrefix/icons/load.gif")
        self.movie.frameChanged.connect(self.repaint)
        self.movie.setScaledSize(self.label_20.size())
        self.label_20.setMovie(self.movie)
        self.movie.start()
        self.label_20.show()

        
    def gif_stop2(self):
        self.pushButton_32.setEnabled(True)
        self.pushButton_46.setEnabled(True)
        self.movie.stop()
        self.label_20.hide()

        
    def read_socket_data2(self):
        global it
        global rec_alarm1
        flash_read = b''
        size = 0
        target_size = 41
        while rec_alarm1:
            con=0
            try:
                while size < target_size:
                        try:    
                            ready_to_read, _, _ = select.select([self.s], [], [], 1.0)
                            if ready_to_read:
                                chunk = self.s.recv(target_size - len(flash_read))
                                if not chunk:
                                    break
                                flash_read += chunk
                                size = len(flash_read)
                            else:
                                rec_alarm1=0
                                size=target_size
                                self.gif_stop2()
                                self.tableWidget_3.setEnabled(True)
                        except socket.error:
                            continue  
            except Exception as e:
                #print(f"Error reading socket: {e}")
                return
            self.process_received_data2(flash_read)
            flash_read=b''
            size=0
            ready_to_read=False
            chunk=b''
    def process_received_data2(self, flash_read1):
        global it
        try:
            QtCore.QCoreApplication.processEvents() 
            d4 = flash_read1.decode()
            datime = d4[3:20]
            print(len(datime))
            if len(datime) == 5:
                QTimer.singleShot(0, self.show_no_data_message)
                QtCore.QCoreApplication.processEvents()
                return   
            f1 = d4[21:25]
            f2 = d4[26:30]
            m1 = d4[31:35]
            m2 = d4[36:40]
            self.tableWidget_3.setItem(it, 0, QTableWidgetItem(datime))
            self.tableWidget_3.setItem(it, 1, QTableWidgetItem(f1))
            self.tableWidget_3.setItem(it, 2, QTableWidgetItem(f2))
            self.tableWidget_3.setItem(it, 3, QTableWidgetItem(m1))
            self.tableWidget_3.setItem(it, 4, QTableWidgetItem(m2))
            QtCore.QCoreApplication.processEvents() 
            it += 1
            self.tableWidget_3.horizontalHeader().setStretchLastSection(True)
            self.tableWidget_3.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
            QtCore.QCoreApplication.processEvents() 
        except Exception as e:
            print(f"Error processing received data: {e}")

            
    def config_get(self):
         global period
         if comport==1:
             QMessageBox.critical(None, "WARNING", f"Change to Ethernet communication")  
         elif comport==2:
             try:
                self.s.send(str.encode(str('CS!')))
                config_data = self.s.recv(1024)
                d2=config_data.decode('ascii')
                serialsize=len(config_data)
                if serialsize==39:
                    interval=d2[3]+d2[4]+d2[5]+d2[6]
                    self.lineEdit_7.setText(interval)
                    time=d2[7]+d2[8]
                    self.lineEdit_8.setText(time)
                    vehicle=d2[9]+d2[10]
                    self.lineEdit_9.setText(vehicle)
                    sigma=d2[11]+d2[12]
                    self.lineEdit_10.setText(sigma)
                    maximum=d2[13]+d2[14]+d2[15]+d2[16]
                    self.lineEdit_11.setText(maximum)
                    minimum=d2[17]+d2[18]+d2[19]+d2[20]
                    self.lineEdit_12.setText(minimum)
                    datime=d2[30]+d2[31]+"/"+d2[33]+d2[34]+"/"+d2[36]+d2[37]+" "+" "+d2[21]+d2[22]+d2[23]+d2[24]+d2[25]+d2[26]+d2[27]+d2[28]
                    self.label.setText(datime)
                    QMessageBox.information(self, "Configuration  get", "Everything is Get!")
                else:
                    QMessageBox.critical(self, "Error", "Failed to Retrieve Configuration Value")
             except:
                 QMessageBox.critical(self, "Error", "Failed to Retrieve Configuration Value")
                 
    def Recorded_clear(self):
        if comport == 1:
            QMessageBox.information(None, "WARNING", "Change to Ethernet communication")
        elif comport == 2:
            self.tableWidget_4.setEnabled(False)
            self.pushButton_46.setEnabled(False)
            self.movie2 = QtGui.QMovie(":/newPrefix/newPrefix/icons/wait.png")
            self.movie2.frameChanged.connect(self.repaint)
            self.movie2.setScaledSize(self.label_440.size())
            self.label_440.setMovie(self.movie2)
            self.movie2.start()
            self.label_440.show()
            QtCore.QCoreApplication.processEvents()
            try:
                self.s.send(str.encode("CM!"))
                self.s.settimeout(10)
                print("Waiting for response...")
                cm_read = self.s.recv(5)
                print(f"Received response: {cm_read}")
                if cm_read == b'CM_1!':
                    self.tableWidget_3.clearContents()
                    self.tableWidget_4.clearContents()
                    self.tableWidget_4.setEnabled(True)
                    self.pushButton_30.setEnabled(True)
                    self.pushButton_46.setEnabled(True)
                    self.gif_stop()
                    QMessageBox.information(None, "DONE", "Extract Alarm data is Erased!")
                else:
                    self.tableWidget_4.setEnabled(True)
                    self.pushButton_30.setEnabled(True)
                    self.pushButton_46.setEnabled(True)
                    self.label_440.hide()
                    QMessageBox.information(None, "ERROR", "Failed to clear recorded data!")
                    
            except socket.timeout:
                self.tableWidget_4.setEnabled(True)
                self.pushButton_30.setEnabled(True)
                self.pushButton_46.setEnabled(True)
                self.label_440.hide()
                QMessageBox.information(None, "ERROR", "No response received within the timeout period.")
                
            except Exception as e:
                self.tableWidget_4.setEnabled(True)
                self.pushButton_30.setEnabled(True)
                self.pushButton_46.setEnabled(True)
                self.label_440.hide()
                QMessageBox.information(None, "ERROR", f"Failed to clear recorded data")

                
    def record_image(self):
        selected_items = self.tableWidget_5.selectedItems()
        if not selected_items:
            QMessageBox.information(None, "Information", f"Select the required row.!")
        if selected_items:
            self.stackedWidget.setCurrentIndex(10)

            
    def datime_view(self):
        current_time = datetime.now()
        x1=current_time.strftime("%H:%M:%S")
        self.label_36.setText(x1)
        self.label_37.setText(x1)
        self.label_38.setText(x1)
        self.label_39.setText(x1)
        self.label_40.setText(x1)
        self.label_41.setText(x1)
        self.label_42.setText(x1)
        self.label_104.setText(x1)
        self.label_48.setText(x1)
        self.label_49.setText(x1)
        self.label_69.setText(x1)
        self.label_71.setText(x1)
        self.label_266.setText(x1)
        self.label_267.setText(x1)
        x=current_time.strftime("%d/%m/%Y")
        self.label_80.setText(x)
        self.label_81.setText(x)
        self.label_82.setText(x)
        self.label_86.setText(x)
        self.label_88.setText(x)
        self.label_90.setText(x)
        self.label_93.setText(x)
        self.label_105.setText(x)
        self.label_106.setText(x)
        self.label_108.setText(x)
        self.label_110.setText(x)
        self.label_112.setText(x)
        self.label_215.setText(x)
        self.label_216.setText(x)
    def Generate_Report(self):
        global main
        main=1
        self.comboBox_5.clear()
        folder_list = [" ","Observed Data","Extracted Alarm Data","Recorded Alarm Data"]
        self.comboBox_5.addItems(folder_list)
        self.lineEdit_14.clear()
        self.lineEdit_13.clear()
        self.tableWidget_5.clear()
        column=[" "," "," "," "," "," "," "," "]
        self.tableWidget_5.setHorizontalHeaderLabels(column)
        header = self.tableWidget_4.horizontalHeader()
        header.setDefaultSectionSize(280)  
        header_font = QtGui.QFont()
        header_font.setFamily("Verdana")
        header_font.setPointSize(17) 
        header_font.setBold(True)     
        header.setFont(header_font)
        self.stackedWidget.setCurrentIndex(11)
    def Recorded_read(self):
        global it
        if comport==1:
             QMessageBox.information(None, "WARNING", f"Change to Ethernet communication")     
        elif comport==2:
             try:
                global alarm_file
                self.tableWidget_4.clearContents()
                self.s.send("FD!".encode())
                current_time = datetime.now()
                alarm_file = current_time.strftime("%d-%m-%Y") + ".csv"
                path = os.getcwd()
                folder_name = "Extracted Alarm Data"
                file = os.path.join(path, folder_name, alarm_file)
                if(os.path.exists(file) and os.path.isfile(file)):
                    os.remove(file) 
                    print("file deleted") 
                else: 
                  print("file not found") 
                global it
                it = 0
                global rec_alarm
                rec_alarm=1
                self.tableWidget_4.setEnabled(False)
                self.gif_exec()
                threading.Thread(target=self.read_socket_data).start()
             except Exception as e:
                QMessageBox.information(None, "WARNING", f"Check the Detector Unit Connectivity")

    def read_socket_data(self):
        global it
        global rec_alarm
        flash_read1 = b''
        size = 0
        target_size = 41
        while rec_alarm:
            con=0
            try:
                while size < target_size:
                        try:    
                            ready_to_read, _, _ = select.select([self.s], [], [], 1.0)
                            if ready_to_read:
                                chunk = self.s.recv(target_size - len(flash_read1))
                                if not chunk:
                                    break
                                flash_read1 += chunk
                                size = len(flash_read1)
                            else:
                                rec_alarm=0
                                size=target_size
                                self.gif_stop()
                                self.tableWidget_4.setEnabled(True)
                        except socket.error:
                            continue  
            except Exception as e:
                #print(f"Error reading socket: {e}")
                return
            self.process_received_data(flash_read1)
            flash_read1=b''
            size=0
            ready_to_read=False
            chunk=b''
    def process_received_data(self, flash_read1):
        global it
        try:
            QtCore.QCoreApplication.processEvents()
            d4 = flash_read1.decode()
            datime = d4[3:20]
            print(len(datime))
            if len(datime) == 5:
                QTimer.singleShot(0, self.show_no_data_message)
                QtCore.QCoreApplication.processEvents()
                return  
            f1 = d4[21:25]
            f2 = d4[26:30]
            m1 = d4[31:35]
            m2 = d4[36:40]
            path = os.getcwd()
            folder_name = "Extracted Alarm Data"
            full_path = os.path.join(path, folder_name, alarm_file)
            header_written = False
            if not os.path.exists(full_path) or os.path.getsize(full_path) == 0:
                header_written = True
            with open(full_path,'a', newline='') as file:
                writer = csv.writer(file)
                if header_written:
                    writer.writerow(['DateTime''F1''F2''M1''M2']) 
                writer.writerow([datime,f1,f2,m1,m2])
            self.tableWidget_4.setItem(it, 0, QTableWidgetItem(datime))
            self.tableWidget_4.setItem(it, 1, QTableWidgetItem(f1))
            self.tableWidget_4.setItem(it, 2, QTableWidgetItem(f2))
            self.tableWidget_4.setItem(it, 3, QTableWidgetItem(m1))
            self.tableWidget_4.setItem(it, 4, QTableWidgetItem(m2))
            it += 1
            self.tableWidget_4.horizontalHeader().setStretchLastSection(True)
            self.tableWidget_4.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        except Exception as e:
            print(f"Error processing received data: {e}")

    def show_no_data_message(self):
        msg = QMessageBox()
        msg.setIcon(QMessageBox.Warning)  # You can change the icon type (e.g., Information, Critical)
        msg.setText("No data received.")
        msg.setWindowTitle("Data Error")
        msg.exec_()  # Show the message box

    def stop_recording(self):
        global rec_alarm
        rec_alarm=0

    def gif_exec(self):
        self.pushButton_30.setEnabled(False)
        self.pushButton_31.setEnabled(False)
        self.pushButton_46.setEnabled(False)
        self.movie = QtGui.QMovie(":/newPrefix/newPrefix/icons/load.gif")
        self.movie.frameChanged.connect(self.repaint)
        self.movie.setScaledSize(self.label_440.size())
        self.label_440.setMovie(self.movie)
        self.movie.start()
        self.label_440.show()
        
    def gif_stop(self):
        self.movie.stop()
        self.label_440.hide()
        self.pushButton_30.setEnabled(True)
        self.pushButton_31.setEnabled(True)
        self.pushButton_46.setEnabled(True)

        
    def gif_exec3(self):
        self.pushButton_30.setEnabled(False)
        self.pushButton_31.setEnabled(False)
        self.pushButton_46.setEnabled(False)
        self.label_440.show()
        print("4")

        
    def gif_stop3(self):
        self.label_440.hide()

        
    def Recorded_Alarm(self):
        if comport==1:
            QMessageBox.information(None, "WARNING", f"Change to Ethernet communication")
        elif comport==2:
            self.clear_socket_buffer()
            self.stackedWidget.setCurrentIndex(4)
            QtCore.QCoreApplication.processEvents()
            
    def open_pdf(self):
        path = os.getcwd()
        file_path =('GUI User Manual_VRCMS.pdf')
        file_path = os.path.abspath(file_path)
        if file_path:
            try:
                if sys.platform == 'win32':
                    os.startfile(file_path)
                elif sys.platform == 'darwin':
                    subprocess.call(('open', file_path))
                else:
                    subprocess.call(('xdg-open', file_path))
            except Exception as e:
                print(f"Error opening file: {e}")

                
    def close_Back(self):
        global host
        global port
        try:
            if(self.serial_port.isOpen()==True):
                self.serial_port.close()
                os._exit(0)
            else:
                self.s.close()
                cv2.destroyAllWindows() 
                os._exit(0)
        except:
            self.s.close()
            print("Ethernet Disconnected")
            cv2.destroyAllWindows() 
            os._exit(0)

    def Back_to_home(self):
        global main
        main==1
        if main==1:
            self.comboBox_5.clear()
            self.lineEdit_13.clear()
            folder_list = [" ","Observed Data","Extracted Alarm Data","Recorded Alarm Data"]
            self.comboBox_5.addItems(folder_list)
            self.pushButton_11.setStyleSheet("\n"
            "color: rgb(176, 176, 148);\n"
            "background-color: rgb(211, 211, 211);\n"
            "border-top: 4px solid black;\n"
            "border-right: 4px solid white;\n"
            "border-left: 4px solid black;\n"
            "border-bottom: 4px solid white;")
            self.pushButton_14.setStyleSheet("\n"
            "color: rgb(176, 176, 148);\n"
            "background-color: rgb(211, 211, 211);\n"
            "border-top: 4px solid black;\n"
            "border-right: 4px solid white;\n"
            "border-left: 4px solid black;\n"
            "border-bottom: 4px solid white;")
            self.pushButton_11.setEnabled(False)
            self.pushButton_14.setEnabled(False)
            self.tableWidget_5.clear()
            self.stackedWidget.setCurrentIndex(1)
        elif main==2:
            #self.stackedWidget.setCurrentIndex(12)
            #QtCore.QCoreApplication.processEvents()
            self.lineEdit_14.clear()
            self.hide()
            self.comboBox_5.clear()
            QtCore.QCoreApplication.processEvents()
            folder_list = [" ","Observed Data","Extracted Alarm Data","Recorded Alarm Data"]
            self.comboBox_5.addItems(folder_list)
            self.pushButton_11.setStyleSheet("\n"
            "color: rgb(176, 176, 148);\n"
            "background-color: rgb(211, 211, 211);\n"
            "border-top: 4px solid black;\n"
            "border-right: 4px solid white;\n"
            "border-left: 4px solid black;\n"
            "border-bottom: 4px solid white;")
            self.pushButton_14.setStyleSheet("\n"
            "color: rgb(176, 176, 148);\n"
            "background-color: rgb(211, 211, 211);\n"
            "border-top: 4px solid black;\n"
            "border-right: 4px solid white;\n"
            "border-left: 4px solid black;\n"
            "border-bottom: 4px solid white;")
            self.pushButton_11.setEnabled(False)
            self.pushButton_14.setEnabled(False)
            self.tableWidget_5.clear()
            column=["  "," "," "," "," "," "," "," "]
            self.tableWidget_5.setHorizontalHeaderLabels(column)
            self.Back_live()

            
    def back_to_report(self):
        self.stackedWidget.setCurrentIndex(11)
        
    def Testmode_Back(self):
        self.label_147.clear()
        self.label_143.clear()
        self.label_148.clear()
        self.label_149.clear()
        self.label_150.clear()
        self.label_144.clear()
        self.label_151.clear()
        self.label_152.clear()
        self.label_145.clear()
        self.label_153.clear()
        self.label_154.clear()
        self.label_155.clear()
        self.stackedWidget.setCurrentIndex(1)

        
    def Home_Back(self):
        self.label_43.clear()
        self.label_44.clear()
        self.label_45.clear()
        self.label_46.clear()
        self.lineEdit_3.clear()
        self.lineEdit_4.clear()
        self.lineEdit_5.clear()
        self.lineEdit_6.clear()
        self.tableWidget_3.clear()
        self.tableWidget_4.clear()
        column=["Date/Time","F1 ","F2","M1","M2"]
        self.tableWidget_3.setHorizontalHeaderLabels(column)
        self.tableWidget_4.setHorizontalHeaderLabels(column)
        header = self.tableWidget_4.horizontalHeader()
        header.setDefaultSectionSize(280)  
        header_font = QtGui.QFont()
        header_font.setFamily("Verdana")
        header_font.setPointSize(17) 
        header_font.setBold(True)     
        header.setFont(header_font)
        header = self.tableWidget_3.horizontalHeader()
        header.setDefaultSectionSize(280)  
        header_font = QtGui.QFont()
        header_font.setFamily("Verdana")
        header_font.setPointSize(17) 
        header_font.setBold(True)     
        header.setFont(header_font)
        self.stackedWidget.setCurrentIndex(1)

        
    def Home_Back1(self):
        self.lineEdit_7.clear()
        self.lineEdit_8.clear()
        self.lineEdit_9.clear()
        self.lineEdit_10.clear()
        self.lineEdit_11.clear()
        self.lineEdit_12.clear()
        self.label.setText("")
        self.stackedWidget.setCurrentIndex(1)

    def IBit(self):
           self.stackedWidget.setCurrentIndex(8)
           self.clear_socket_buffer()
           QtCore.QCoreApplication.processEvents()
           try:
               self.s.send(str.encode(str('IB!')))
               try:
                    self.s.settimeout(0.01)  
                    while True:
                        dataibit = self.s.recv(1024)
                        print(dataibit)
                        if not dataibit:
                            break
                        else:
                           serialsize=len(dataibit)
                           d11=dataibit.decode('ascii')
                           if serialsize==8:
                                if d11[3]=='0' :
                                    self.label_250.setStyleSheet("background-color: rgb(218, 0, 0);\n"
                "border-top: 6px solid black;\n"
                "border-right: 6px solid white;\n"
                "border-left: 6px solid black;\n"
                "border-bottom: 6px solid white;")
                                else:
                                    self.label_250.setStyleSheet("background-color: rgb(62, 167, 46);\n"
                "border-top: 6px solid black;\n"
                "border-right: 6px solid white;\n"
                "border-left: 6px solid black;\n"
                "border-bottom: 6px solid white;")
                                if d11[4]=='0':
                                    self.label_252.setStyleSheet("background-color: rgb(218, 0, 0);\n"
                "border-top: 6px solid black;\n"
                "border-right: 6px solid white;\n"
                "border-left: 6px solid black;\n"
                "border-bottom: 6px solid white;")
                                else:
                                    self.label_252.setStyleSheet("background-color: rgb(62, 167, 46);\n"
                "border-top: 6px solid black;\n"
                "border-right: 6px solid white;\n"
                "border-left: 6px solid black;\n"
                "border-bottom: 6px solid white;")
                                if d11[5]=='0':
                                    self.label_253.setStyleSheet("background-color: rgb(218, 0, 0);\n"
                "border-top: 6px solid black;\n"
                "border-right: 6px solid white;\n"
                "border-left: 6px solid black;\n"
                "border-bottom: 6px solid white;") 
                                else:
                                    self.label_253.setStyleSheet("background-color: rgb(62, 167, 46);\n"
                "border-top: 6px solid black;\n"
                "border-right: 6px solid white;\n"
                "border-left: 6px solid black;\n"
                "border-bottom: 6px solid white;")  
                                if d11[6]=='0':
                                    self.label_254.setStyleSheet("background-color: rgb(218, 0, 0);\n"
                "border-top: 6px solid black;\n"
                "border-right: 6px solid white;\n"
                "border-left: 6px solid black;\n"
                "border-bottom: 6px solid white;")  
                                else:
                                    self.label_254.setStyleSheet("background-color: rgb(62, 167, 46);\n"
                                    "border-top: 6px solid black;\n"
                                    "border-right: 6px solid white;\n"
                                    "border-left: 6px solid black;\n"
                                    "border-bottom: 6px solid white;")
               except socket.timeout:
                   pass  
               except socket.error as e:
                   print(f"Socket error: {e}")  
               finally:
                   self.s.settimeout(None)
            
           except Exception as e:
               print(f"Error: {e}")
               self.stackedWidget.setCurrentIndex(8)
    def PBit(self):
        self.stackedWidget.setCurrentIndex(7)
        self.clear_socket_buffer()
        QtCore.QCoreApplication.processEvents()
        try:
           self.s.send(str.encode(str('PB!')))
           try:
               self.s.settimeout(0.01)  
               while True:
                   data_pbit = self.s.recv(1024)
                   if not data_pbit:
                       break
                   else:
                       serialsize=len(data_pbit)
                       d11=data_pbit.decode('ascii')
                       if serialsize==8:
                            if d11[3]=='0' :
                                self.label_236.setStyleSheet("background-color: rgb(218, 0, 0);\n"
            "border-top: 6px solid black;\n"
            "border-right: 6px solid white;\n"
            "border-left: 6px solid black;\n"
            "border-bottom: 6px solid white;")
                            else:
                                self.label_236.setStyleSheet("background-color: rgb(62, 167, 46);\n"
            "border-top: 6px solid black;\n"
            "border-right: 6px solid white;\n"
            "border-left: 6px solid black;\n"
            "border-bottom: 6px solid white;")
                            if d11[4]=='0':
                                self.label_237.setStyleSheet("background-color: rgb(218, 0, 0);\n"
            "border-top: 6px solid black;\n"
            "border-right: 6px solid white;\n"
            "border-left: 6px solid black;\n"
            "border-bottom: 6px solid white;")
                            else:
                                self.label_237.setStyleSheet("background-color: rgb(62, 167, 46);\n"
            "border-top: 6px solid black;\n"
            "border-right: 6px solid white;\n"
            "border-left: 6px solid black;\n"
            "border-bottom: 6px solid white;")
                            if d11[5]=='0':
                                self.label_238.setStyleSheet("background-color: rgb(218, 0, 0);\n"
            "border-top: 6px solid black;\n"
            "border-right: 6px solid white;\n"
            "border-left: 6px solid black;\n"
            "border-bottom: 6px solid white;")
                            else:
                                self.label_238.setStyleSheet("background-color: rgb(62, 167, 46);\n"
            "border-top: 6px solid black;\n"
            "border-right: 6px solid white;\n"
            "border-left: 6px solid black;\n"
            "border-bottom: 6px solid white;")
                            if d11[6]=='0':
                                self.label_239.setStyleSheet("background-color: rgb(218, 0, 0);\n"
            "border-top: 6px solid black;\n"
            "border-right: 6px solid white;\n"
            "border-left: 6px solid black;\n"
            "border-bottom: 6px solid white;")
                            else:
                                self.label_239.setStyleSheet("background-color: rgb(62, 167, 46);\n"
            "border-top: 6px solid black;\n"
            "border-right: 6px solid white;\n"
            "border-left: 6px solid black;\n"
            "border-bottom: 6px solid white;")
           except socket.timeout:
                pass  
           except socket.error as e:
                print(f"Socket error: {e}")  
           finally:
                self.s.settimeout(None)
                
        except Exception as e:
            print(f"Error: {e}")
            #self.stackedWidget.setCurrentIndex(7)

    def Bit(self):
        self.stackedWidget.setCurrentIndex(6)
        labels = [
            self.label_236, self.label_237, self.label_238,
            self.label_239, self.label_250, self.label_252,
            self.label_253, self.label_254
        ]
        
        style = """
        background-color: rgb(218, 0, 0);
        border-top: 6px solid black;
        border-right: 6px solid white;
        border-left: 6px solid black;
        border-bottom: 6px solid white;
        """
        
        for label in labels:
            label.setStyleSheet(style)

    def Serial_connect(self):
        global comport
        comport=1
        self.pushButton_12.setStyleSheet("color: rgb(255, 255, 255);\n"
"color: rgb(90, 255, 87);\n"
"background-color: rgb(129, 255, 75);\n"
"color: rgb(0, 0, 0);\n"
"\n"
"border-top: 6px solid black;\n"
"border-right: 6px solid white;\n"
"border-left: 6px solid black;\n"
"border-bottom:6px  solid white;")
        self.pushButton_13.setStyleSheet("color: rgb(255, 255, 255);\n"
"color: rgb(90, 255, 87);\n"
"background-color: rgb(216, 216, 216);\n"
"color: rgb(0, 0, 0);\n"
"\n"
"border-top: 6px solid black;\n"
"border-right: 6px solid white;\n"
"border-left: 6px solid black;\n"
"border-bottom:6px  solid white;")

        
    def Ethernet_connect(self):
        global comport
        comport=2
        self.pushButton_13.setStyleSheet("color: rgb(255, 255, 255);\n"
"color: rgb(90, 255, 87);\n"
"background-color: rgb(129, 255, 75);\n"
"color: rgb(0, 0, 0);\n"
"\n"
"border-top: 6px solid black;\n"
"border-right: 6px solid white;\n"
"border-left: 6px solid black;\n"
"border-bottom:6px  solid white;")
        self.pushButton_12.setStyleSheet("color: rgb(255, 255, 255);\n"
"color: rgb(90, 255, 87);\n"
"background-color: rgb(216, 216, 216);\n"
"color: rgb(0, 0, 0);\n"
"\n"
"border-top: 6px solid black;\n"
"border-right: 6px solid white;\n"
"border-left: 6px solid black;\n"
"border-bottom:6px  solid white;")

        
    def port_refresh(self):
        self.comboBox_2.clear()
        self.serial(serial)

        
    def communication_port(self):
        self.comboBox_2.clear()
        ports = serial.tools.list_ports.comports()
        for port in ports:
            self.comboBox_2.addItem(port.device)
        self.stackedWidget.setCurrentIndex(5)

        
    def alarm_set(self):
        if comport==1:
           QMessageBox.critical(None, "WARNING", f"Change Ethernet communication")  
        elif comport==2:
            try:
                self.s.send(str.encode(str("UA_"+self.lineEdit_3.text()+"@"+self.lineEdit_4.text()+"@"+self.lineEdit_5.text()+"@"+self.lineEdit_6.text()+"!")))
                setalarm1=self.s.recv(5)
                print(setalarm1)
                if setalarm1==b'UA_1!':
                   QMessageBox.information(None, "Done","Alarm Level is Set")
                else:
                   QMessageBox.critical(None, "Done","Alarm Level Setting is FAiled")
            except Exception as e:
                QMessageBox.critical(self, "Error", "Alarm setting failed. Please check the console for details.")

                
    def alarm_get(self):
        if comport==1:
             QMessageBox.critical(None, "WARNING", f"Change to Ethernet communication")  
        elif comport==2:
            try:
                self.s.send(str.encode(str('UA!')))
                alarm_data= self.s.recv(4096)
                serialsize=len(alarm_data)
                if serialsize==43:
                    d3=alarm_data.decode('ascii')
                    bgf1=d3[3]+d3[4]+d3[5]+d3[6]
                    self.label_43.setText(bgf1)
                    bgf2=d3[8]+d3[9]+d3[10]+d3[11]
                    self.label_44.setText(bgf2)
                    bgf3=d3[13]+d3[14]+d3[15]+d3[16]
                    self.label_45.setText(bgf3)
                    bgf4=d3[18]+d3[19]+d3[20]+d3[21]
                    self.label_46.setText(bgf4)
                    alarmf1=d3[23]+d3[24]+d3[25]+d3[26]
                    self.lineEdit_3.setText(alarmf1)
                    alarmf2=d3[28]+d3[29]+d3[30]+d3[31]
                    self.lineEdit_4.setText(alarmf2)
                    alarmf3=d3[33]+d3[34]+d3[35]+d3[36]
                    self.lineEdit_5.setText(alarmf3)
                    alarmf4=d3[38]+d3[39]+d3[40]+d3[41]
                    self.lineEdit_6.setText(alarmf4)
                    QMessageBox.information(self, "Alarm get", "Everything is Get!")
                else:
                    QMessageBox.critical(self, "Error", "Failed to Retrieve Alarm Value") 
            except:
                 QMessageBox.critical(self, "Error", "Failed to Retrieve Alarm Value")

                 
    def config_set(self):
        global config_datetime
        current_time=datetime.now()
        config_datetime = current_time.strftime('%H:%M:%S:%d:%m:%y')
        if comport==1:
            QMessageBox.information(None, "WARNING", f"Change to Ethernet communication")  
        elif comport==2:
            try:
                self.s.send(str.encode(str("CS_"+self.lineEdit_7.text()+"@"+self.lineEdit_8.text()+"@"+self.lineEdit_9.text()+"@"+self.lineEdit_10.text()+"@"+self.lineEdit_11.text()+"@"+self.lineEdit_12.text()+"@"+config_datetime+"!")))
                Interval=self.s.recv(5)
                if Interval==b'CS_1!':
                    QMessageBox.information(None, "Done", f"Configuration is set!")
                else:
                    QMessageBox.information(None, "Configuration setting is Failed")
            except Exception as e:
                QMessageBox.critical(self, "Error", "Configuration Setting Failed. Please Check the Console for Details.")

    def Alarm_level(self):
        try:
            if comport==1:
                 QMessageBox.critical(None, "WARNING", f"Change to Ethernet communication")             
            elif comport==2:
                self.stackedWidget.setCurrentIndex(3)
                self.clear_socket_buffer()
                #QtCore.QCoreApplication.processEvents()
                self.s.send(str.encode(str('UA!')))
                try:
                    self.s.settimeout(0.1)  
                    while True:
                        alarm_data = self.s.recv(1024)
                        if not alarm_data:
                            break
                        else:
                            serialsize=len(alarm_data) 
                            if serialsize==43:
                                d3=alarm_data.decode('ascii')
                                QtCore.QCoreApplication.processEvents()
                                bgf1=d3[3]+d3[4]+d3[5]+d3[6]
                                self.label_43.setText(bgf1)
                                bgf2=d3[8]+d3[9]+d3[10]+d3[11]
                                self.label_44.setText(bgf2)
                                bgf3=d3[13]+d3[14]+d3[15]+d3[16]
                                self.label_45.setText(bgf3)
                                bgf4=d3[18]+d3[19]+d3[20]+d3[21]
                                self.label_46.setText(bgf4)
                                alarmf1=d3[23]+d3[24]+d3[25]+d3[26]
                                self.lineEdit_3.setText(alarmf1)
                                alarmf2=d3[28]+d3[29]+d3[30]+d3[31]
                                self.lineEdit_4.setText(alarmf2)
                                alarmf3=d3[33]+d3[34]+d3[35]+d3[36]
                                self.lineEdit_5.setText(alarmf3)
                                alarmf4=d3[38]+d3[39]+d3[40]+d3[41]
                                self.lineEdit_6.setText(alarmf4)
                except socket.timeout:
                    pass  
                except socket.error as e:
                    print(f"Socket error: {e}")  
                finally:
                    self.s.settimeout(None)
        except Exception as e:
             print(f"Error: {e}")

    def clear_socket_buffer1(self):
        try:
            self.s.settimeout(0.1)  
            while True:
                chunk = self.s.recv(1024)
                if not chunk:
                    break  
        except socket.timeout:
            pass  
        except socket.error as e:
            print(f"Socket error: {e}")  
        finally:
            self.s.settimeout(None)
    def clear_socket_buffer(self):
        try:
            self.s.settimeout(0.1)  
            while True:
                data = self.s.recv(1024)
                if not data:
                    break  
        except socket.timeout:
            pass  
        except socket.error as e:
            print(f"Socket error: {e}")  
        finally:
            self.s.settimeout(None)
            
    def configuration(self):
        try:
            if comport==1:
                QMessageBox.critical(None, "WARNING", f"Change to Ethernet communication")
            elif comport==2:
                self.stackedWidget.setCurrentIndex(2)
                self.clear_socket_buffer()
                #QtCore.QCoreApplication.processEvents()
                self.s.send(str.encode(str('CS!')))
                try:
                    self.s.settimeout(0.1)  
                    while True:
                        config_data = self.s.recv(1024)
                        if not config_data:
                            break
                        else:
                            d2=config_data.decode('ascii')
                            serialsize=len(config_data)
                            if serialsize==39:
                                interval=d2[3]+d2[4]+d2[5]+d2[6]
                                self.lineEdit_7.setText(interval)
                                time=d2[7]+d2[8]
                                self.lineEdit_8.setText(time)
                                vehicle=d2[9]+d2[10]
                                self.lineEdit_9.setText(vehicle)
                                sigma=d2[11]+d2[12]
                                self.lineEdit_10.setText(sigma)
                                maximum=d2[13]+d2[14]+d2[15]+d2[16]
                                self.lineEdit_11.setText(maximum)
                                minimum=d2[17]+d2[18]+d2[19]+d2[20]
                                self.lineEdit_12.setText(minimum)
                                datime=d2[30]+d2[31]+"/"+d2[33]+d2[34]+"/"+d2[36]+d2[37]+" "+" "+d2[21]+d2[22]+d2[23]+d2[24]+d2[25]+d2[26]+d2[27]+d2[28]
                                self.label.setText(datime)
                            elif serialsize=="":
                                dialog = QMessageBox()
                                dialog.setText("No Response!")
                                dialog.exec_()
                except socket.timeout:
                    pass  
                except socket.error as e:
                    print(f"Socket error: {e}")  
                finally:
                    self.s.settimeout(None)
            
        except Exception as e:
            print(f"Exception: {e}")
    
    def keyPressEvent(self, event):
        if event.key() in {Qt.Key_Return, Qt.Key_Enter}:
            current_page = self.stackedWidget.currentIndex() 

            if current_page == 0:
                self.login() 
            elif current_page == 13:
                self.live_settings() 
        
        super().keyPressEvent(event)
    def login(self):
        global comport, host, port
        user = "Vrcms"
        password = "Vrcms123"
        user1 = self.lineEdit.text()
        password1 = self.lineEdit_2.text()
        if user== user1 and password == password1:
            global newfile, newfile1, newfile2
            current_time = datetime.now()
            newfile = current_time.strftime("%d-%m-%Y") + ".csv"
            newfile1 = current_time.strftime("%d-%m-%Y") + ".csv"
            newfile2 = current_time.strftime("%d-%m-%Y") + ".csv"
            directories = ["2024", "Recorded Alarm Data", "Observed Data", "Extracted Alarm Data", "Screenshot"]
            for directory in directories:
                if not os.path.exists(directory):
                    os.makedirs(directory)
            if comport == 1:
                QMessageBox.information(None, "WARNING", f"Change to Ethernet communication")
            elif comport == 2:
                try:
                    if hasattr(self, 's') and self.s:
                        self.s.close() 
                    self.s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
                    self.s.settimeout(3)
                    self.s.connect((host, port))
                    print("Ethernet connected")
                    self.label_224.setText("ETHERNET")
                    self.label_224.setStyleSheet("background-color: rgb(62, 167, 46);\n"
                        "border-top: 6px solid black;\n"
                        "border-right: 6px solid white;\n"
                        "border-left: 6px solid black;\n"
                        "border-bottom: 6px solid white;"
                    )
                    self.stackedWidget.setCurrentIndex(12)
                    self.s.send("LD1!".encode())
                    self.timer9 = QtCore.QTimer()
                    self.timer9.timeout.connect(self.read_live_data)
                    self.timer9.start(1000)
                except socket.error as e:
                    QMessageBox.critical(self, "ERROR", f"Failed to connect to Ethernet: {e}")
                    self.timer9 = QtCore.QTimer()
                    self.timer9.timeout.connect(self.read_live_data)
                    self.timer9.start(1000)
                    self.label_224.setText("ETHERNET")
                    self.label_224.setStyleSheet("background-color: rgb(218,0,0);\n"
                        "border-top: 6px solid black;\n"
                        "border-right: 6px solid white;\n"
                        "border-left: 6px solid black;\n"
                        "border-bottom: 6px solid white;"
                    )
                    self.stackedWidget.setCurrentIndex(12)
        else:
            QMessageBox.critical(self, "Error", "User ID or Password is Invalid.")

    def login_view(self):
         if self.pushButton_4.isChecked():
            self.lineEdit_2.setEchoMode(QtWidgets.QLineEdit.Normal)
            self.pushButton_4.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/visible.png);")
         else:
             self.pushButton_4.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/un_visible.png);")
             self.lineEdit_2.setEchoMode(QtWidgets.QLineEdit.Password)
    def sublogin_view(self):
         if self.pushButton_19.isChecked():
            self.lineEdit_14.setEchoMode(QtWidgets.QLineEdit.Normal)
            self.pushButton_19.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/visible.png);")
         else:
             self.pushButton_19.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/un_visible.png);")
             self.lineEdit_14.setEchoMode(QtWidgets.QLineEdit.Password)
    def sblogin_cancel(self):
        self.Back_live()
        #self.stackedWidget.setCurrentIndex(12)
    def cancel(self):
        cv2.destroyAllWindows()
        os._exit(0)
    def sub_login(self):
        self.lineEdit_13.clear()
        self.lineEdit_14.clear()
        self.timer9.stop()
        self.stackedWidget.setCurrentIndex(13)
    def live_settings(self):
        try:
            user = "Admin"
            password = "Vrcms123"
            user1 = self.lineEdit_13.text()
            password1 = self.lineEdit_14.text()
            if user== user1 and password == password1:
                    self.stackedWidget.setCurrentIndex(1)
                    QtCore.QCoreApplication.processEvents()
                    self.s.send(str.encode("LD0!"))
                    self.timer9.stop()
                   # threading.Thread(target=self.port_connect4).start()
            else:
                 QMessageBox.critical(self, "ERROR", f"login attempt failed")
            #QtCore.QCoreApplication.processEvents()
           # self.s.send(str.encode("LD0!"))
            #self.timer9.stop()
        except:
            print("ADS")
            self.timer9.stop()
    def live_report(self):
        global main
        main=2
        self.stackedWidget.setCurrentIndex(11)
        QtCore.QCoreApplication.processEvents()
        self.comboBox_5.clear()
        self.tableWidget_5.clear()
        column=["  "," "," "," "," "," "," "," "]
        self.tableWidget_5.setHorizontalHeaderLabels(column)
        self.pushButton_11.setStyleSheet("\n"
"color: rgb(176, 176, 148);\n"
"background-color: rgb(211, 211, 211);\n"
"border-top: 4px solid black;\n"
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
        self.pushButton_14.setStyleSheet("\n"
"color: rgb(176, 176, 148);\n"
"background-color: rgb(211, 211, 211);\n"
"border-top: 4px solid black;\n"
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
        folder_list = [" ","Observed Data","Extracted Alarm Data","Recorded Alarm Data"]
        self.comboBox_5.addItems(folder_list)
        self.lineEdit_14.clear()
        self.lineEdit_13.clear()
        self.tableWidget_5.clear()
        column=[" "," "," "," "," "," "," "," "]
        self.tableWidget_5.setHorizontalHeaderLabels(column)
        header = self.tableWidget_4.horizontalHeader()
        header.setDefaultSectionSize(280)  
        header_font = QtGui.QFont()
        header_font.setFamily("Verdana")
        header_font.setPointSize(17) 
        header_font.setBold(True)     
        header.setFont(header_font)
        try:
            self.s.send(str.encode(str('LD0!')))
            self.timer9.stop()
            self.tableWidget_5.clear()
            column=[" "," "," "," "," "," "," "," "]
            self.tableWidget_5.setHorizontalHeaderLabels(column)
            self.pushButton_11.setStyleSheet("\n"
"color: rgb(176, 176, 148);\n"
"background-color: rgb(211, 211, 211);\n"
"\n"
"border-top: 4px solid black;\n"
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
            self.pushButton_14.setStyleSheet("\n"
"color: rgb(176, 176, 148);\n"
"background-color: rgb(211, 211, 211);\n"
"border-top: 4px solid black;\n"
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
        except:
            self.timer9.stop()
    def msg_close(self):
        msg_box = QMessageBox()
        msg_box.setWindowTitle("Information")
        msg_box.setText("Application close!")
        #cancel_button = msg_box.addButton("Cancel", QMessageBox.RejectRole)
        close_button = msg_box.addButton("Close", QMessageBox.AcceptRole)
        msg_box.setDefaultButton(close_button)
        result = msg_box.exec_()
        if msg_box.clickedButton() == close_button:
            os._exit(0)
            print("Close button clicked")
        else:
            print("not click")
    def Quit(self):
        global host
        global port
        try:
            if(self.serial_port.isOpen()==True):
                self.serial_port.close()
                cv2.destroyAllWindows()
                os._exit(0)
            else:
                self.timer9.stop()
                self.s.close()
                cv2.destroyAllWindows()
                os._exit(0)
        except:
            self.s.close()
            cv2.destroyAllWindows()
            os._exit(0)
    def read_live_data(self):
        self.s.setblocking(False)
        global it,lop_count
        chunk=b''
        flash_read = b''
        size = 0
        target_size = 30
        timeout = 0.1
        if hasattr(self, 's') and self.s:
            ready_to_read, ready_to_write, in_error = select.select([self.s], [], [], timeout)
            if ready_to_read:
                try:
                    while size < target_size:
                        chunk = self.s.recv(target_size - size)
                        self.clear_socket_buffer()
                        if not chunk:
                            break
                        flash_read += chunk
                        size += len(chunk)
                    if size == target_size:
                        lop_count = 0
                        self.liveData1(flash_read)
                except socket.error as e:
                    pass
                    lop_count += 1
                    if lop_count > 4:
                        self.hide()
                        #threading.Thread(target=self.port_connect3).start()
                        self.port_connect3()
            else:
                lop_count += 1
                if lop_count > 4:
                    self.hide()
                    #threading.Thread(target=self.port_connect3).start()
                    self.port_connect3()
        else:
            print("Socket is not initialized.")
    def hide(self):
        self.label_208.hide()
        self.label_210.hide()
        self.label_209.hide()
        self.label_211.hide()
        self.label_312.hide()
        self.pushButton_6.hide()
        self.label_61.clear()
        self.label_58.clear()
        self.label_59.clear()
        self.label_60.clear()
        self.pushButton_90.setEnabled(True)
        self.pushButton_9.setEnabled(True)
        self.pushButton_10.setEnabled(True)
        self.label_15.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/VRCMS_IMG.jpg);\n")
        self.pushButton_90.setStyleSheet(" border-top: 4px solid black;\n"
"border-right:  4px  solid white;\n"
" border-left:  4px  solid black;\n"
"border-bottom: 4px  solid white;\n"
"color: rgb(0,0,0);\n"                                    
"background-color: rgb(180, 180, 180);")
        self.label_61.setStyleSheet(" border-top: 4px solid black;\n"
"border-right:  4px  solid white;\n"
" border-left:  4px  solid black;\n"
"border-bottom: 4px  solid white;\n"
"background-color: rgb(180, 180, 180);")
        self.label_60.setStyleSheet(" border-top:  4px  solid black;\n"
"border-right:  4px  solid white;\n"
" border-left:  4px  solid black;\n"
"border-bottom: 4px  solid white;\n"
"background-color: rgb(180, 180, 180);")
        self.label_58.setStyleSheet(" border-top:  4px  solid black;\n"
"border-right:  4px  solid white;\n"
" border-left:  4px  solid black;\n"
"border-bottom: 4px  solid white;\n"
"background-color: rgb(180, 180, 180);")
        self.label_59.setStyleSheet(" border-top:  4px  solid black;\n"
"border-right:  4px  solid white;\n"
" border-left:  4px  solid black;\n"
"border-bottom: 4px  solid white;\n"
"background-color: rgb(180, 180, 180);")
        self.label_224.setStyleSheet(
"background-color: rgb(218,0,0);\n"
"border-top: 6px solid black;\n"
"border-right: 6px solid white;\n"
"border-left: 6px solid black;\n"
"border-bottom: 6px solid white;")
        self.label_16.setStyleSheet(" border-top: 4px  solid black;\n"
"border-right:  4px  solid white;\n"
" border-left:  4px  solid black;\n"
"border-bottom: 4px  solid white;\n"
"background-color: rgb(199, 199, 199);")
        self.label_17.setStyleSheet(" border-top: 4px  solid black;\n"
"border-right: 4px  solid white;\n"
" border-left:  4px  solid black;\n"
"border-bottom: 4px  solid white;\n"
"background-color: rgb(199, 199, 199);")
        self.label_18.setStyleSheet(" border-top: 4px  solid black;\n"
"border-right: 4px  solid white;\n"
" border-left:  4px  solid black;\n"
"border-bottom: 4px  solid white;\n"
"background-color: rgb(199, 199, 199);")
        self.label_19.setStyleSheet(" border-top: 4px  solid black;\n"
"border-right: 4px  solid white;\n"
" border-left:  4px  solid black;\n"
"border-bottom: 4px  solid white;\n"
"background-color: rgb(199, 199, 199);")
        self.label_66.setStyleSheet(" border-top: 4px  solid black;\n"
"border-right: 4px  solid white;\n"
" border-left:  4px  solid black;\n"
"border-bottom: 4px  solid white;\n"
"background-color: rgb(199, 199, 199);")
        self.label_15.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/VRCMS_IMG.jpg);\n"
"\n"
"\n"
"")
        if self.thread is not None and self.thread.is_alive():
                    self.stop_event.set()
                    self.thread.join()
        if self.thread is not None:
                            self.stop_thread_event.set()  
                            self.thread.join()
                            self.thread = None
                            self.stop_thread_event.clear()
    def liveData1(self,flash_read):
        QtCore.QCoreApplication.processEvents()
        p = 0
        global code
        global image_name
        global image_name2
        global day_name
        global newfile1
        global newfile
        global newfile2
        global image_path
        global image_path1
        global x
        global vec_cont
        global obsrv_count
        global recorded_count
        global vehicle
        current_date = datetime.now().strftime('%d-%m-%Y')
        current_time = datetime.now().strftime('%H:%M:%S')
        period=0
        livecmt=flash_read
        serialsize=len(livecmt)
        if serialsize==30:
            QtCore.QCoreApplication.processEvents()
            self.label_224.setStyleSheet(
                "background-color: rgb(62, 167, 46);\n"
                "border-top: 6px solid black;\n"
                "border-right: 6px solid white;\n"
                "border-left: 6px solid black;\n"
                "border-bottom: 6px solid white;"
            )
            d3=livecmt.decode()
            if d3[3]=='0':
                self.pushButton_9.setEnabled(True)
                self.pushButton_10.setEnabled(True)
                self.pushButton_90.setEnabled(True)
                QtCore.QCoreApplication.processEvents()
                self.pushButton_6.hide()
                self.label_58.clear()
                self.label_59.clear()
                self.label_60.clear()
                self.label_61.clear()
                self.label_209.show()
                self.label_210.hide()
                self.label_211.hide()
                self.label_208.hide()
                self.label_312.hide()
                self.label_209.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/arrow.png);")
                self.label_15.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/VRCMS_IMG.jpg);")
                self.label_16.setStyleSheet("border-top: 4px solid black;\n"
"background-color: rgb(104, 231, 250);\n"
"border-bottom: 4px solid white;\n"
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"")
                self.label_17.setStyleSheet("border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
                self.label_18.setStyleSheet("border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
                self.label_19.setStyleSheet("border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
                self.label_66.setStyleSheet("border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
                self.label_58.setStyleSheet("border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
                self.label_59.setStyleSheet(" border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
                self.label_60.setStyleSheet(" border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
                self.label_61.setStyleSheet(" border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
                self.pushButton_90.setStyleSheet(" border-top: 4px solid black;\n"
"background-color: rgb(180,180,180);\n"                                           
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
            elif d3[3]=='1':
                self.pushButton_6.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/unmute.png);")
                if self.thread is not None and self.thread.is_alive():
                    self.stop_event.set()
                    self.thread.join() 
                QtCore.QCoreApplication.processEvents()
                self.pushButton_9.setEnabled(True)
                self.pushButton_10.setEnabled(True)
                self.pushButton_90.setEnabled(True)
                vec_cont=1
                recorded_count=1
                obsrv_count=1
                x=1
                self.pic = 0
                self.label_58.clear()
                self.label_59.clear()
                self.label_60.clear()
                self.label_61.clear()
                self.pushButton_6.hide()
                self.label_210.show()
                self.label_209.hide()
                self.label_211.hide()
                self.label_208.hide()
                self.label_312.hide()
                self.label_210.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/arrow.png);")
                self.label_15.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/Ready.png);")
                self.label_17.setStyleSheet("border-top: 4px solid black;\n"
"background-color: rgb(104, 231, 250);\n"
"border-bottom: 4px solid white;\n"
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
        "")
                self.label_16.setStyleSheet("border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-bottom: 4px solid white;\n"
"border-right: 4px solid white;\n"
"border-left: 4px solid black;")
                self.label_18.setStyleSheet("border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-bottom: 4px solid white;\n"
"border-right: 4px solid white;\n"
"border-left: 4px solid black;")
                self.label_19.setStyleSheet("border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-bottom: 4px solid white;\n"
"border-right: 4px solid white;\n"
"border-left: 4px solid black;")
                self.label_66.setStyleSheet("border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-bottom: 4px solid white;\n"
"border-right: 4px solid white;\n"
"border-left: 4px solid black;")
                self.label_58.setStyleSheet("border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
                self.label_59.setStyleSheet(" border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
                self.label_60.setStyleSheet(" border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
                self.label_61.setStyleSheet(" border-top: 4px solid black;\n"
"background-color: rgb(199, 199, 199);\n"                                           
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
                self.pushButton_90.setStyleSheet(" border-top: 4px solid black;\n"
"background-color: rgb(180,180,180);\n"                                           
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
                if self.thread is not None:
                    self.stop_thread_event.set()  
                    self.thread.join()
                    self.thread = None
                    self.stop_thread_event.clear()
            elif d3[3]=='2':#image
                 QtCore.QCoreApplication.processEvents()
                 self.pushButton_9.setEnabled(False)
                 self.pushButton_10.setEnabled(False)
                 self.pushButton_90.setEnabled(False)
                 try:
                    self.pushButton_6.hide()
                    self.label_58.clear()
                    self.label_59.clear()
                    self.label_60.clear()
                    self.label_61.clear()
                    self.label_211.show()
                    self.label_210.hide()
                    self.label_209.hide()
                    self.label_208.hide()
                    self.label_312.hide()
                    self.label_211.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/arrow.png);")
                    self.label_15.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/monitor.png);")
                    self.label_18.setStyleSheet("border-top: 4px solid black;\n"
    "background-color: rgb(104, 231, 250);\n"
    "border-bottom: 4px  solid white;\n"
    "border-right:4px solid white;\n"
    "border-left: 4px solid black;\n""")
                    self.label_16.setStyleSheet("border-top: 4px solid black;\n"
    "background-color: rgb(199, 199, 199);\n"                                           
    "border-right: 4px solid white;\n"
    "border-left: 4px solid black;\n"
    "border-bottom: 4px solid white;")
                    self.label_17.setStyleSheet("border-top: 4px solid black;\n"
    "background-color: rgb(199, 199, 199);\n"                                            
    "border-right: 4px solid white;\n"
    "border-left: 4px solid black;\n"
    "border-bottom: 4px solid white;")
                    self.label_19.setStyleSheet("border-top: 4px solid black;\n"
    "background-color: rgb(199, 199, 199);\n"                                            
    "border-right: 4px solid white;\n"
    "border-left: 4px solid black;\n"
    "border-bottom: 4px solid white;")
                    self.label_66.setStyleSheet("border-top:4px solid black;\n"
    "background-color: rgb(199, 199, 199);\n" 
    "border-right: 4px solid white;\n"
    "border-left: 4px solid black;\n"
    "border-bottom: 4px solid white;")
                    QtCore.QCoreApplication.processEvents()
                    global a
                    global my_thread
                    print(vec_cont)
                    if vec_cont==1:
                        current_date = datetime.now().strftime('%d-%m-%Y')
                        count = 0
                        try:
                            with open('vehicle_count.txt', 'r') as file:
                                lines = file.readlines()
                                if lines:
                                    last_line = lines[-1].strip()
                                    parts = last_line.split(' - ')
                                    if len(parts) > 1:
                                        last_count = int(parts[0])
                                        last_date = parts[1]
                                        if last_date == current_date:
                                            count = last_count
                                        else:
                                            count = 0
                                    else:
                                        count = 0
                        except FileNotFoundError:
                            count = 0
                        except ValueError:
                            count = 0
                        count += 1
                        try:
                            with open('vehicle_count.txt', 'w') as file:
                                file.write(f"{count} - {current_date}\n")
                        except IOError as e:
                            print(f"Error writing to file: {e}")
                        self.label_67.setText(f"{count}")
                        QtCore.QCoreApplication.processEvents()
                        current_time = datetime.now()
                        cap = cv2.VideoCapture(url1)
                        cap2 = cv2.VideoCapture(url2)
                        success,image=cap.read()
                        print(success)
                        success1,image1=cap2.read()
                        print(success1)
                        image_name = current_time.strftime("%d-%m-%Y_%H-%M-%S-cam1")
                        image_name2 = current_time.strftime("%d-%m-%Y_%H-%M-%S-cam2")
                        year_name = "2024"
                        path = os.getcwd()
                        path1 = os.path.join(path ,year_name)
                        if os.path.exists(year_name):
                             try:
                                if success and success1:
                                    year_name = "2024"
                                    day_name = current_time.strftime("%d-%m-%Y")
                                    path = os.getcwd()
                                    path1 = os.path.join(path, year_name)
                                    path2 = os.path.join(path1, day_name)
                                    print(path2)
                                    if not os.path.exists(path2):
                                        os.makedirs(path2)
                                    image_name = current_time.strftime("%d-%m-%Y_%H-%M-%S-cam1")
                                    image_name2 = current_time.strftime("%d-%m-%Y_%H-%M-%S-cam2")
                                    image_path = os.path.join(path2, image_name + ".jpg")
                                    image_path1 = os.path.join(path2, image_name2 + ".jpg")
                                    cv2.imwrite(image_path, image)
                                    cv2.imwrite(image_path1, image1)
                                    vec_cont=0
                                    print("Images saved successfully.")
                                else:
                                    print("Failed to capture frames from one or both cameras.")
                             except:
                                 self.pixmap = QPixmap(':/newPrefix/newPrefix/icons/disconnect.png')
                                 self.label_6.setPixmap(self.pixmap)
                                 self.label_6.show()
                                 self.label_21.setPixmap(self.pixmap)
                                 self.label_21.show()   
                        else:
                            year_folder = current_time.strftime("%Y")
                            year_name = os.makedirs(year_folder)
                    else:
                        print(vec_cont)
                 except:
                    print("error")
            elif d3[3]=='3':#clear
                self.pushButton_9.setEnabled(False)
                self.pushButton_10.setEnabled(False)
                self.pushButton_90.setEnabled(False)
                QtCore.QCoreApplication.processEvents()
                f1=d3[10]+d3[11]+d3[12]+d3[13]
                self.label_59.setText(f1) 
                f2=d3[15]+d3[16]+d3[17]+d3[18]
                self.label_58.setText(f2)
                m1=d3[20]+d3[21]+d3[22]+d3[23]
                self.label_61.setText(m1)
                m2=d3[25]+d3[26]+d3[27]+d3[28]
                self.label_60.setText(m2)
                count3 = self.label_67.text()
                if obsrv_count==1:
                     path = os.getcwd()
                     folder_name2 = "Observed Data"
                     file_path = os.path.join(path, folder_name2, newfile2)
                     header_written = False
                     if not os.path.exists(file_path) or os.path.getsize(file_path) == 0:
                        header_written = True
                     with open(file_path, 'a', newline='') as file:
                            writer = csv.writer(file)
                            if header_written:
                                writer.writerow(['Date','Time','Count', 'F1', 'F2', 'M1', 'M2']) 
                            writer.writerow([current_date,current_time,count3, f1, f2, m1, m2])
                     file.close()
                     obsrv_count=0
                else:
                    print(obsrv_count)
                self.label_58.setStyleSheet("background-color: rgb(62, 167, 46);")
                self.label_59.setStyleSheet("background-color: rgb(62, 167, 46);")
                self.label_60.setStyleSheet("background-color: rgb(62, 167, 46);")
                self.label_61.setStyleSheet("background-color: rgb(62, 167, 46);")
                self.pushButton_6.hide()
                self.label_312.show()
                self.label_210.hide()
                self.label_209.hide()
                self.label_208.hide()
                self.label_211.hide()
                self.label_312.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/arrow.png);")
                self.label_15.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/clear1.png);")
                self.label_66.setStyleSheet("border-top: 4px solid black;\n"
    "background-color: rgb(104, 231, 250);\n"
    "border-bottom: 4px solid white;\n"
    "border-right: 4px solid white;\n"
    "border-left: 4px solid black;\n"
    "")
                self.label_16.setStyleSheet("border-top: 4px solid black;\n"
    "background-color: rgb(199, 199, 199);\n"                                           
    "border-right: 4px solid white;\n"
    "border-left: 4px solid black;\n"
    "border-bottom: 4px solid white;")
                self.label_17.setStyleSheet("border-top: 4px solid black;\n"
    "background-color: rgb(199, 199, 199);\n"                                            
    "border-right:4px solid white;\n"
    "border-left: 4px solid black;\n"
    "border-bottom: 4px solid white;")
                self.label_19.setStyleSheet("border-top: 4px  solid black;\n"
    "background-color: rgb(199, 199, 199);\n"                                            
    "border-right: 4px solid white;\n"
    "border-left: 4px solid black;\n"
    "border-bottom: 4px solid white;")
                self.label_18.setStyleSheet("border-top: 4px solid black;\n"
    "background-color: rgb(199, 199, 199);\n" 
    "border-right: 4px solid white;\n"
    "border-left: 4px solid black;\n"
    "border-bottom: 4px solid white;")
                try:
                    x=1
                    year_name = "2024"
                    path = os.getcwd()
                    path1 = os.path.join(path , year_name)
                    if os.path.exists(year_name):
                        path2 = os.path.join(path1 ,day_name)
                        removing_files = glob.glob(os.path.join(path2, f"{image_name}.jpg")) + \
                                         glob.glob(os.path.join(path2, f"{image_name2}.jpg"))
                        for file_path in removing_files:
                            os.remove(file_path)
                            #print(f"Removed: {file_path}")
                    else:
                        print(f"Error: Directory '{path1}' does not exist.")
                except Exception as e:
                    print(f"Error: {e}")
            elif d3[3]=='4' and self.pic == 0:#contaminated
                 if self.thread is None or not self.thread.is_alive():
                    self.stop_event.clear()
                    self.thread = threading.Thread(target=self.my_thread_func)
                    self.thread.start()
                 self.pushButton_9.setEnabled(False)
                 self.pushButton_10.setEnabled(False)
                 self.pushButton_90.setEnabled(False)
                 self.pushButton_90.setStyleSheet("\n"
"color: rgb(176, 176, 148);\n"
"background-color: rgb(211, 211, 211);\n"
"\n"
"border-top: 4px solid black;\n"
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
                 QtCore.QCoreApplication.processEvents()
                 f1=d3[10]+d3[11]+d3[12]+d3[13]
                 self.label_59.setText(f1)
                 f2=d3[15]+d3[16]+d3[17]+d3[18]
                 self.label_58.setText(f2)
                 m1=d3[20]+d3[21]+d3[22]+d3[23]
                 self.label_61.setText(m1)
                 m2=d3[25]+d3[26]+d3[27]+d3[28]
                 self.label_60.setText(m2)
                 count3 = self.label_67.text()
                 path = os.getcwd()
                 if obsrv_count==1:#obsrv data_saved
                     folder_name2 = "Observed Data"
                     file_path = os.path.join(path, folder_name2, newfile2)
                     header_written = False
                     if not os.path.exists(file_path) or os.path.getsize(file_path) == 0:
                        header_written = True
                     with open(file_path, 'a', newline='') as file:
                            writer = csv.writer(file)
                            if header_written:
                                writer.writerow(['Date','Time', 'Count', 'F1', 'F2', 'M1', 'M2']) 
                            writer.writerow([current_date,current_time, count3, f1, f2, m1, m2])
                     file.close()
                     obsrv_count=0
                 else:
                    print(obsrv_count)
                 if recorded_count==1:#Recorded data_saved
                    folder_name ="Recorded Alarm Data"
                    file_path1 = os.path.join(path, folder_name, newfile1)
                    header_written = False
                    if not os.path.exists(file_path1) or os.path.getsize(file_path1) == 0:
                        header_written = True
                    with open(file_path1, 'a', newline='') as file:
                          writer = csv.writer(file)
                          if header_written:
                              writer.writerow(['Date','Time', 'Count', 'F1', 'F2', 'M1', 'M2','CAMERA-01','CAMERA-02'])
                          writer.writerow([(current_date),(current_time),(count3),(f1),(f2),(m1),(m2),(image_path),(image_path1)])
                    file.close()
                    recorded_count=0
                 else:
                    print(recorded_count)
                 self.label_16.setStyleSheet(" border-top: 3px solid black;\n"
    "border-right: 3px solid white;\n"
    " border-left: 3px solid black;\n"
    "border-bottom:3px solid white;\n"
    "background-color: rgb(199, 199, 199);")
                 self.label_17.setStyleSheet(" border-top: 3px solid black;\n"
    "border-right: 3px solid white;\n"
    " border-left: 3px solid black;\n"
    "border-bottom:3px solid white;\n"
    "background-color: rgb(199, 199, 199);")
                 self.label_18.setStyleSheet(" border-top: 3px solid black;\n"
    "border-right: 3px solid white;\n"
    " border-left: 3px solid black;\n"
    "border-bottom:3px solid white;\n"
    "background-color: rgb(199, 199, 199);")
                 self.label_211.hide()
                 self.label_15.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/monitor.png);")
                 #self.label_15.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/clear1.png);")
                 self.label_208.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/arrow.png);")
                 self.pushButton_6.show()
                 self.label_208.show()
                 self.label_19.setStyleSheet(" border-top: 3px solid black;\n"
"border-right: 3px solid white;\n"
" border-left: 3px solid black;\n"
"border-bottom:3px solid white;\n"
"background-color: rgb(218, 0, 0);")
                 self.label_66.setStyleSheet(" border-top: 3px solid black;\n"
    "border-right: 3px solid white;\n"
    " border-left: 3px solid black;\n"
    "border-bottom:3px solid white;\n"
    "background-color: rgb(199, 199, 199);")
                 QtCore.QCoreApplication.processEvents()
                 if d3[5]=='0':
                    self.label_61.setStyleSheet("background-color: rgb(62, 167, 46);")
                 else:
                    self.label_61.setStyleSheet("background-color: rgb(255, 0, 0);")
                 if d3[6]=='0':
                    self.label_60.setStyleSheet("background-color: rgb(62, 167, 46);")
                 else:
                    self.label_60.setStyleSheet("background-color: rgb(255, 0, 0);")
                 if d3[7]=='0':
                    self.label_59.setStyleSheet("background-color: rgb(62, 167, 46);")
                 else:
                    self.label_59.setStyleSheet("background-color: rgb(255, 0, 0);")
                 if d3[8]=='0':
                    self.label_58.setStyleSheet("background-color: rgb(62, 167, 46);")
                 else :
                    self.label_58.setStyleSheet("background-color: rgb(255, 0, 0);")
                 QtCore.QCoreApplication.processEvents()
                 self.label_210.hide()
                 self.label_209.hide()
                 self.label_312.hide()
                 self.pic += 1
                 folder_name = "screenshot"
                 path = os.getcwd()
                 path1 = os.path.join(path,folder_name)
                 current_time = datetime.now()
                 if not os.path.exists(path1):
                    os.makedirs(path1)
                 day_name = current_time.strftime("%d-%m-%Y")
                 picture_path = os.path.join(path1,day_name)
                 if not os.path.exists(picture_path):
                    os.makedirs(picture_path)
                 screenshot = QApplication.primaryScreen().grabWindow(QApplication.desktop().winId())
                 pixmap = QPixmap(screenshot)
                 screen_time = datetime.now().strftime("%H-%M-%S")
                 file_name = os.path.join(picture_path, f"screenshot_{screen_time}.png")
                 if pixmap.save(file_name, "PNG"):
                     pass
                    #print(f"Screenshot saved as {file_name}")     
                 else:
                    print("Failed to save screenshot.")
            elif d3[3]=='4' and self.pic != 0:#contaminated
                self.pushButton_90.setStyleSheet("\n"
"color: rgb(176, 176, 148);\n"
"background-color: rgb(211, 211, 211);\n"
"border-top: 4px solid black;\n"
"border-right: 4px solid white;\n"
"border-left: 4px solid black;\n"
"border-bottom: 4px solid white;")
                self.label_16.setStyleSheet(" border-top: 3px solid black;\n"
"border-right: 3px solid white;\n"
" border-left: 3px solid black;\n"
"border-bottom:3px solid white;\n"
"background-color: rgb(199, 199, 199);")
                self.label_17.setStyleSheet(" border-top: 3px solid black;\n"
"border-right: 3px solid white;\n"
" border-left: 3px solid black;\n"
"border-bottom:3px solid white;\n"
"background-color: rgb(199, 199, 199);")
                self.label_18.setStyleSheet(" border-top: 3px solid black;\n"
"border-right: 3px solid white;\n"
" border-left: 3px solid black;\n"
"border-bottom:3px solid white;\n"
"background-color: rgb(199, 199, 199);")
                self.label_19.setStyleSheet(" border-top: 3px solid black;\n"
"border-right: 3px solid white;\n"
" border-left: 3px solid black;\n"
"border-bottom:3px solid white;\n"
"background-color: rgb(218, 0, 0);")
                self.label_66.setStyleSheet(" border-top: 3px solid black;\n"
"border-right: 3px solid white;\n"
" border-left: 3px solid black;\n"
"border-bottom:3px solid white;\n"
"background-color: rgb(199, 199, 199);")
                self.label_15.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/monitor.png);")
               # self.label_15.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/clear1.png);")
                self.label_208.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/arrow.png);")
                self.pushButton_6.show()
                self.label_208.show()
                self.label_210.hide()
                self.label_209.hide()
                self.label_211.hide()
                self.label_312.hide()
                self.label_61.setStyleSheet("background-color: rgb(255, 0, 0);" if d3[5] == '0' else "background-color: rgb(62, 167, 46);")
                self.label_60.setStyleSheet("background-color: rgb(255, 0, 0);" if d3[6] == '0' else "background-color: rgb(62, 167, 46);")
                self.label_59.setStyleSheet("background-color: rgb(255, 0, 0);" if d3[7] == '0' else "background-color: rgb(62, 167, 46);")
                self.label_58.setStyleSheet("background-color: rgb(255, 0, 0);" if d3[8] == '0' else "background-color: rgb(62, 167, 46);")        
            if code==0:
                x=1
                if d3[5]=='0':
                    self.label_61.setStyleSheet("background-color: rgb(62, 167, 46);")
                else:
                    self.label_61.setStyleSheet("background-color: rgb(255, 0, 0);")
                if d3[6]=='0':
                    self.label_60.setStyleSheet("background-color: rgb(62, 167, 46);")
                else:
                    self.label_60.setStyleSheet("background-color: rgb(255, 0, 0);")
                if d3[7]=='0':
                    self.label_59.setStyleSheet("background-color: rgb(62, 167, 46);")
                else:
                    self.label_59.setStyleSheet("background-color: rgb(255, 0, 0);")
                if d3[8]=='0':
                    self.label_58.setStyleSheet("background-color: rgb(62, 167, 46);")
                else :
                    self.label_58.setStyleSheet("background-color: rgb(255, 0, 0);")
            elif code==1:
                x=1
                if d3[5]=='0':
                    self.label_61.setStyleSheet("background-color: rgb(62, 167, 46);")
                else:
                    self.label_61.setStyleSheet("background-color: rgb(255, 0, 0);")
                if d3[6]=='0':
                    self.label_60.setStyleSheet("background-color: rgb(62, 167, 46);")
                else:
                    self.label_60.setStyleSheet("background-color: rgb(255, 0, 0);")
                if d3[7]=='0':
                    self.label_59.setStyleSheet("background-color: rgb(62, 167, 46);")
                else:
                    self.label_59.setStyleSheet("background-color: rgb(255, 0, 0);")
                if d3[8]=='0':
                    self.label_58.setStyleSheet("background-color: rgb(62, 167, 46);")
                else :
                    self.label_58.setStyleSheet("background-color: rgb(255, 0, 0);")

    def my_thread_func(self):
        audio_path = 'Audio/beep1.mp3'
        audio_path = os.path.abspath(audio_path)
     
        while not self.stop_event.is_set():
            try:
                playsound(audio_path)
            except Exception as e:
                print(f"Error playing sound: {e}")
            time.sleep(0.05)
        pass
       # print("Thread has been stopped")
    def horn(self):
        if self.thread is None or not self.thread.is_alive():
            self.pushButton_6.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/unmute.png);")
            self.stop_event.clear()
            self.thread = threading.Thread(target=self.my_thread_func)
            self.thread.start()
        else:
            self.stop_event.set() 
            self.pushButton_6.setStyleSheet("border-image: url(:/newPrefix/newPrefix/icons/mute.png);")

    def display_frame(self, q_img):
         #print("g")
        self.label_21.setPixmap(q_img)
        # Ensure the label is shown
        self.label_21.show()

    def show_warning(self):
        # Show the warning image if no video feed
        self.pixmap = QPixmap(':/newPrefix/newPrefix/icons/disconnect.png')
        self.label_21.setPixmap(self.pixmap)
        self.label_21.show()

    def display_frame2(self, q_img):
        #print("h")
        self.label_6.setPixmap(q_img)
        # Ensure the label is shown
        self.label_6.show()

    def show_warning2(self):
        # Show the warning image if no video feed
        self.pixmap = QPixmap(':/newPrefix/newPrefix/icons/disconnect.png')
        self.label_6.setPixmap(self.pixmap)
        self.label_6.show()

    def closeEvent(self, event):
        # Stop the capture thread on close
        self.capture_thread1.stop()
        self.capture_thread2.stop()
        self.port_thread.stop()
        event.accept()
        print("c")
if __name__ == "__main__":
    import sys
    app = QtWidgets.QApplication(sys.argv)
    mainWindow = MainWindow()
    mainWindow.showMaximized()
    mainWindow.show()
    sys.exit(app.exec_())
